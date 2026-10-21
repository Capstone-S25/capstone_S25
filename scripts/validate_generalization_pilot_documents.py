"""Check the new report/deck identities, numeric text, layout and rendering."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import xml.etree.ElementTree as ET
from pathlib import Path

from pptx import Presentation

from scripts.validate_progress_presentation_20261019 import render_contact_sheets

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "reports/generalization_pilot_20261020"
SLIDES = ROOT / "reports/slides"
STEM = "model_generalization_pilot_20261020"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def norm(s):
    return re.sub(r"\s+", " ", s).strip()


def check_pdf(path):
    info = subprocess.check_output(["pdfinfo", str(path)], text=True)
    pages = int(re.search(r"Pages:\s+(\d+)", info)[1])
    tree = ET.fromstring(
        subprocess.check_output(["pdftotext", "-bbox", str(path), "-"], text=True)
    )
    ns = {"x": "http://www.w3.org/1999/xhtml"}
    words = 0
    for i, page in enumerate(tree.findall(".//x:page", ns), 1):
        w, h = float(page.attrib["width"]), float(page.attrib["height"])
        for word in page.findall(".//x:word", ns):
            a = word.attrib
            assert 5 <= float(a["xMin"]) < float(a["xMax"]) <= w - 5, (
                path,
                i,
                word.text,
            )
            assert 5 <= float(a["yMin"]) < float(a["yMax"]) <= h - 5, (
                path,
                i,
                word.text,
            )
            words += 1
    fonts = subprocess.check_output(["pdffonts", str(path)], text=True).splitlines()[2:]
    assert fonts and all(
        re.search(r"\byes\s+(?:yes|no)\s+(?:yes|no)\s+\d+\s+\d+$", s) for s in fonts
    )
    return {"pages": pages, "words_within_page_bounds": words, "fonts_embedded": True}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--render", action="store_true")
    args = parser.parse_args()
    report_manifest = json.loads((REPORT / "source_manifest.json").read_text())
    for path, sha in report_manifest["sources_sha256"].items():
        assert digest(ROOT / path) == sha, path
    for path, sha in report_manifest["documents_sha256"].items():
        assert digest(REPORT / path) == sha, path
    for path, sha in report_manifest["exports_sha256"].items():
        assert digest(REPORT / path) == sha, path
    manifest = json.loads((SLIDES / (STEM + "_manifest.json")).read_text())
    for path, sha in manifest["sources_sha256"].items():
        assert digest(ROOT / path) == sha, path
    assert (
        digest(SLIDES / "build_generalization_pilot_20261020.py")
        == manifest["builder_sha256"]
    )
    presentation = Presentation(SLIDES / (STEM + ".pptx"))
    assert len(presentation.slides) == manifest["slides"] == 10
    for slide, plan in zip(presentation.slides, manifest["layout"]):
        text = norm(" ".join(s.text for s in slide.shapes if s.has_text_frame))
        for item in plan["items"]:
            if item["kind"] == "text":
                assert norm(item["value"]) in text, (plan["title"], item["value"])
                assert (
                    len(item["value"].splitlines()) * item["leading"]
                    <= item["h"] + 0.01
                )
        assert norm(plan["notes"]) in norm(slide.notes_slide.notes_text_frame.text)
    result = {
        "status": "passed",
        "report": check_pdf(REPORT / "runthrough.pdf"),
        "slides": check_pdf(SLIDES / (STEM + ".pdf")),
        "pptx_slides_and_notes_checked": 10,
        "report_source_files_checked": len(report_manifest["sources_sha256"]),
        "no_model_inference": True,
    }
    assert result["slides"]["pages"] == 10
    log = (REPORT / "runthrough.log").read_text()
    assert not any(
        s in log for s in ("Missing character", "Overfull", "undefined references")
    )
    for path, tokens in (
        (
            REPORT / "runthrough.pdf",
            ("0.7309", "0.6227", "0.7768", "0.3855", "0.3573", "52.42", "weak"),
        ),
        (
            SLIDES / (STEM + ".pdf"),
            ("0.7309", "0.6227", "0.7768", "0.3855", "0.3573", "proxy"),
        ),
    ):
        text = subprocess.check_output(
            ["pdftotext", "-layout", str(path), "-"], text=True
        )
        assert all(token in text for token in tokens), path
    if args.render:
        for path, directory in (
            (REPORT / "runthrough.pdf", REPORT / "previews"),
            (SLIDES / (STEM + ".pdf"), SLIDES / "previews_pilot_20261020"),
        ):
            pages, sheets = render_contact_sheets(path, directory)
            (directory / "index.md").write_text(
                "# Rendered review\n\n"
                + "\n\n".join(f"![Pages]({s})" for s in sheets)
                + "\n"
            )
            assert pages == check_pdf(path)["pages"]
    result["artifacts_sha256"] = {
        str(p.relative_to(ROOT)): digest(p)
        for p in (
            REPORT / "runthrough.pdf",
            REPORT / "runthrough.tex",
            REPORT / "report.md",
            SLIDES / (STEM + ".pdf"),
            SLIDES / (STEM + ".pptx"),
            Path(__file__),
        )
    }
    (REPORT / "document_validation.json").write_text(
        json.dumps(result, indent=2) + "\n"
    )
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
