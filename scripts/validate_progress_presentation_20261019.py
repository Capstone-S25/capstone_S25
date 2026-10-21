"""Check the updated PDF/PPTX and optionally render every presentation page."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports/slides"
STEM = "model_evaluation_progress_20261019"
SUMMARY = "summary_slide_progress_20261019"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def render_contact_sheets(pdf, destination, *, columns=3, rows=3, thumb_width=480):
    """Render all pages with Poppler and retain compact, numbered contact sheets."""
    destination.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="capstone_pdf_review_") as temporary:
        prefix = Path(temporary) / "page"
        subprocess.run(
            [
                "pdftoppm",
                "-scale-to",
                "1100",
                "-jpeg",
                "-jpegopt",
                "quality=85",
                str(pdf),
                str(prefix),
            ],
            check=True,
            stdout=subprocess.DEVNULL,
        )
        paths = sorted(Path(temporary).glob("page-*.jpg"))
        first = Image.open(paths[0])
        thumb_height = round(thumb_width * first.height / first.width)
        cell_height = thumb_height + 28
        sheets, index = [], []
        per_sheet = columns * rows
        for start in range(0, len(paths), per_sheet):
            sheet = Image.new(
                "RGB",
                (columns * (thumb_width + 12) + 12, rows * (cell_height + 12) + 12),
                "#e6ebed",
            )
            draw = ImageDraw.Draw(sheet)
            name = f"overview_{len(sheets) + 1}.jpg"
            for offset, path in enumerate(paths[start : start + per_sheet]):
                x = 12 + (offset % columns) * (thumb_width + 12)
                y = 12 + (offset // columns) * (cell_height + 12)
                with Image.open(path) as page:
                    page.thumbnail(
                        (thumb_width, thumb_height), Image.Resampling.LANCZOS
                    )
                    sheet.paste(page, (x, y + 24))
                draw.text((x + 4, y + 5), f"Page {start + offset + 1}", fill="#142c40")
                index.append({"page": start + offset + 1, "contact_sheet": name})
            sheet.save(destination / name, quality=92)
            sheets.append(name)
        (destination / "page_index.json").write_text(json.dumps(index, indent=2) + "\n")
        return len(paths), sheets


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--render", action="store_true")
    args = parser.parse_args()
    manifest_path = OUT / "progress_source_manifest_20261019.json"
    manifest = json.loads(manifest_path.read_text())
    for path, expected in manifest["sources_sha256"].items():
        assert digest(ROOT / path) == expected, path
    for path, expected in manifest["builders_sha256"].items():
        assert digest(OUT / path) == expected, path
    assert manifest["slide_count"] == 45
    ppt = Presentation(OUT / (STEM + ".pptx"))
    assert len(ppt.slides) == 45
    notes = 0
    for actual, planned in zip(ppt.slides, manifest["layout"]):
        text = "\n".join(shape.text for shape in actual.shapes if shape.has_text_frame)
        normalized = lambda s: re.sub(r"\s+", " ", s).strip()
        for item in planned["items"]:
            if item["kind"] == "text":
                assert normalized(item["value"]) in normalized(text), (
                    planned["title"],
                    item["value"],
                )
                assert (
                    len(item["value"].splitlines()) * item["leading"]
                    <= item["h"] + 0.01
                )
        assert normalized(planned["notes"]) in normalized(
            actual.notes_slide.notes_text_frame.text
        )
        notes += 1
    pages = 0
    for stem, expected in [(STEM, 45), (SUMMARY, 1)]:
        pdf = OUT / (stem + ".pdf")
        info = subprocess.check_output(["pdfinfo", str(pdf)], text=True)
        assert int(re.search(r"Pages:\s+(\d+)", info)[1]) == expected
        assert len(Presentation(OUT / (stem + ".pptx")).slides) == expected
        tree = ET.fromstring(
            subprocess.check_output(["pdftotext", "-bbox", str(pdf), "-"], text=True)
        )
        ns = {"x": "http://www.w3.org/1999/xhtml"}
        for i, page in enumerate(tree.findall(".//x:page", ns), 1):
            w, h = float(page.attrib["width"]), float(page.attrib["height"])
            for word in page.findall(".//x:word", ns):
                a = word.attrib
                assert 5 <= float(a["xMin"]) < float(a["xMax"]) <= w - 5, (
                    stem,
                    i,
                    word.text,
                )
                assert 5 <= float(a["yMin"]) < float(a["yMax"]) <= h - 5, (
                    stem,
                    i,
                    word.text,
                )
        fonts = subprocess.check_output(["pdffonts", str(pdf)], text=True).splitlines()[
            2:
        ]
        assert fonts and all(
            re.search(r"\byes\s+(?:yes|no)\s+(?:yes|no)\s+\d+\s+\d+$", v) for v in fonts
        )
        pages += expected
    text = subprocess.check_output(
        ["pdftotext", "-layout", str(OUT / (STEM + ".pdf")), "-"], text=True
    )
    for token in ["70.8%", "0.1822", "0.6933", "4,341,760", "187", "retrospective"]:
        assert token in text, token
    destination = OUT / "previews_20261019"
    sheets = []
    if args.render:
        rendered, sheets = render_contact_sheets(OUT / (STEM + ".pdf"), destination)
        assert rendered == 45
        with tempfile.TemporaryDirectory(prefix="capstone_summary_") as temporary:
            prefix = Path(temporary) / "summary"
            subprocess.run(
                [
                    "pdftoppm",
                    "-scale-to",
                    "1920",
                    "-singlefile",
                    "-png",
                    str(OUT / (SUMMARY + ".pdf")),
                    str(prefix),
                ],
                check=True,
            )
            (OUT / (SUMMARY + ".png")).write_bytes(
                prefix.with_suffix(".png").read_bytes()
            )
        (destination / "index.md").write_text(
            "# Updated progress deck — all rendered pages\n\n"
            + "\n".join(
                f"- [Slides {i * 9 + 1}–{min(45, (i + 1) * 9)}]({name})."
                for i, name in enumerate(sheets)
            )
            + "\n"
        )
    result = {
        "status": "complete",
        "slides": 45,
        "editable_pptx_slides": 45,
        "speaker_notes": notes,
        "all_source_and_builder_hashes_match": True,
        "pdf_pages_checked_including_summary": pages,
        "fonts_embedded": True,
        "pdf_text_within_page_boundaries": True,
        "text_fit_and_geometry": "checked by builder; rechecked against saved layout",
        "primary_artifact_sha256": {
            stem + suffix: digest(OUT / (stem + suffix))
            for stem in [STEM, SUMMARY]
            for suffix in [".pdf", ".pptx"]
        },
        "source_manifest_sha256": digest(manifest_path),
    }
    if args.render:
        result.update(pdf_pages_rendered=45, contact_sheets=sheets)
    review = destination / "review.json"
    if review.exists():
        reviewed = json.loads(review.read_text())
        assert reviewed["pdf_sha256"] == digest(OUT / (STEM + ".pdf"))
        assert reviewed["pages_rendered"] == 45
        result["pdf_pages_rendered"] = reviewed["pages_rendered"]
        result["visual_review"] = reviewed
    (OUT / "progress_validation_20261019.json").write_text(
        json.dumps(result, indent=2) + "\n"
    )
    print(
        f"Validated 45 editable slides and notes, 46 PDF pages, source identities and embedded fonts; rendered={args.render}."
    )


if __name__ == "__main__":
    main()
