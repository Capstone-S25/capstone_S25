"""Independently check exported stage metrics and the four-model presentation."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import subprocess
import xml.etree.ElementTree as ET
from functools import cache
from pathlib import Path

import numpy as np
from pptx import Presentation

from scripts.validate_progress_presentation_20261019 import render_contact_sheets

ROOT = Path(__file__).resolve().parents[1]
SLIDES = ROOT / "reports/slides"
DETAIL = SLIDES / "training_evaluation_20261020"
STEM = "model_training_evaluation_20261020"


@cache
def load(path):
    return json.loads((ROOT / path).read_text())


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def expected(row):
    study, model, stage, task, source, point = (
        row[k] for k in ("study", "model", "stage", "task", "source", "operating_point")
    )
    seed = int(row["seed"])
    if study == "clip_visual_15oct":
        data = load(
            f"results/clip_finetuning_20261015/seed_{seed}/{model}/float32_evaluation/results.json"
        )
        cell = data[stage]["external" if source == "stanford40" else "test"][task]
        return cell["operating_points"][point], cell["thresholds"][point]
    if study == "strategy_16oct":
        data = load(
            f"results/strategy_study_20261016/{model}/transfer_evaluation/per_source_results.json"
        )
        key = next(k for k in data if k.endswith("/" + task + "/" + source))
        selected = load(f"results/strategy_study_20261016/{model}/selection.json")[
            "tasks"
        ][task]
        return data[key]["selected"][point], selected["thresholds"][point]
    if study == "smol_lora_19oct":
        selected = load(
            f"results/smolvlm_finetuning_20261019/seed_{seed}/selection.json"
        )
        route = (
            "original"
            if stage == "original"
            else selected["joint_adapter"]
            if stage == "joint"
            else selected["task_routes"][task]
        )
        name = route.replace("/", "__")
        if source in (
            "hmdb51_action_transfer",
            "commons_new_author_transfer",
            "meva_abandon_action_transfer",
        ):
            data = load(
                f"results/smolvlm_finetuning_20261019/additional_transfer/evaluation/{name}/results.json"
            )[task]
        else:
            split = "external" if source == "stanford40" else "test"
            data = load(
                f"results/smolvlm_finetuning_20261019/evaluation/{name}/results.json"
            )[split][task]
        return data["sources"][source][point], data["thresholds"][point]
    if study == "vadclip_session_13oct":
        if stage == "temporal_adapter":
            cell = load("results/gpu_all_models_20261013/vadclip/adapter/eval.json")[
                "metrics"
            ]["dual_mean"]
        else:
            key = "zeroshot" if stage == "zero_shot" else "linearprobe"
            cell = load(
                "results/gpu_all_models_20261013/vadclip/baselines/baselines.json"
            )["metrics"][key]
        cell = {
            **cell,
            "ap": cell["aupr"],
            "fpr": cell["false_positive_rate"],
            "positive": cell["n_pos"],
            "negative": cell["n_neg"],
        }
        return cell, cell["threshold"]
    if study == "video_pilot_16oct":
        data = load(f"results/video_pilot_20261016/{model}/results.json")[
            "clip_metrics"
        ][task][stage][source]
        selection = load(f"results/video_pilot_20261016/{model}/joint_selection.json")[
            "tasks"
        ][task]
        candidate = next(c for c in selection["candidates"] if c["name"] == stage)
        return data["operating_points"][point], candidate["thresholds"][point]
    if study == "source_policy_17oct":
        cell = load("results/event_improvements_20261017/smoking_results.json")[model][
            stage
        ]
        return cell["sources"][source], cell["threshold"]
    if study == "source_ablation_20oct":
        cell = load(
            f"results/generalization_pilot_20261020/smoking/{model}/results.json"
        )["results"][stage + "/" + source]
        selection = load(
            f"results/generalization_pilot_20261020/smoking/{model}/selection.json"
        )
        return cell["operating_points"][point], selection["candidates"][
            cell["candidate"]
        ]["thresholds"][point]
    if study == "source_transfer_20oct":
        kind = point.removesuffix("_budget")
        cell = load(
            f"results/generalization_pilot_20261020/smolvlm_transfer/{stage}/results.json"
        )[kind]["dev_fpr_budget"]
        key = (
            "new_source_thresholds"
            if kind == "new_source_calibrated"
            else "original19oct_thresholds"
        )
        calibration = load(
            f"results/generalization_pilot_20261020/smolvlm_transfer/{stage}/calibration.json"
        )
        return cell, calibration[key]["dev_fpr_budget"]
    raise AssertionError(study)


def metrics():
    records = list(csv.DictReader((DETAIL / "metrics.csv").open()))
    for row in records:
        cell, threshold = expected(row)
        assert row["threshold"] and np.isclose(
            float(row["threshold"]), threshold, atol=1e-12, rtol=1e-12
        ), row
        for key in (
            "n",
            "positive",
            "negative",
            "tp",
            "fp",
            "fn",
            "tn",
            "ap",
            "auroc",
            "precision",
            "recall",
            "f1",
            "fpr",
        ):
            assert np.isclose(float(row[key]), cell[key], atol=1e-12, rtol=1e-12), (
                row,
                key,
            )
        assert int(row["tp"]) + int(row["fn"]) == int(row["positive"])
        assert int(row["fp"]) + int(row["tn"]) == int(row["negative"])
        accuracy = (cell["tp"] + cell["tn"]) / cell["n"]
        balanced = (cell["recall"] + 1 - cell["fpr"]) / 2
        assert np.isclose(float(row["accuracy"]), accuracy)
        assert np.isclose(float(row["balanced_accuracy"]), balanced)
    return len(records)


def datasets():
    records = list(csv.DictReader((DETAIL / "dataset_counts.csv").open()))
    names = {
        "clip_visual_15oct": "data/clip_training_20261015/samples.jsonl",
        "expanded_strategy_smol_19oct": "results/smolvlm_finetuning_20261019/samples.jsonl",
        "additional_transfer_19oct": "results/smolvlm_finetuning_20261019/additional_transfer/samples.jsonl",
        "xrf55_proxy_20oct": "data/generalization_pilot_20261020/xrf55/samples.jsonl",
    }
    manifests = {
        k: [json.loads(s) for s in (ROOT / path).read_text().splitlines()]
        for k, path in names.items()
    }
    for row in records:
        if row["study"] == "vadclip_sessions_13oct":
            split = load("data/models/vadclip_features_cuda/manifest.json")[
                row["split"]
            ]
            source = load(
                "results/gpu_all_models_20261013/vadclip/source_validation.json"
            )["sessions"]
            n = sum(source[s]["n_frames"] for s in split)
            positive = sum(source[s]["n_positive"] for s in split)
        else:
            group = [
                r
                for r in manifests[row["study"]]
                if r["task"] == row["task"] and r["split"] == row["split"]
            ]
            n = len(group)
            positive = sum(r["target"] for r in group)
        assert (int(row["n"]), int(row["positive"]), int(row["negative"])) == (
            n,
            positive,
            n - positive,
        ), row
    return len(records)


def normalized(s):
    return re.sub(r"\s+", " ", s).strip()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--render", action="store_true")
    args = parser.parse_args()
    manifest = json.loads((DETAIL / "manifest.json").read_text())
    for key in ("sources_sha256", "builders_sha256"):
        for name, sha in manifest[key].items():
            assert digest(ROOT / name) == sha, name
    metric_count = metrics()
    dataset_count = datasets()
    assert (
        metric_count == manifest["metric_rows"]
        and dataset_count == manifest["dataset_rows"]
    )
    presentation = Presentation(SLIDES / (STEM + ".pptx"))
    assert len(presentation.slides) == manifest["slide_count"]
    for slide, planned in zip(presentation.slides, manifest["layout"]):
        text = normalized(
            " ".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        )
        for item in planned["items"]:
            if item["kind"] == "text":
                assert normalized(item["value"]) in text, (
                    planned["title"],
                    item["value"],
                )
                assert (
                    len(item["value"].splitlines()) * item["leading"]
                    <= item["h"] + 0.01
                )
        assert normalized(planned["notes"]) in normalized(
            slide.notes_slide.notes_text_frame.text
        )
    pdf = SLIDES / (STEM + ".pdf")
    info = subprocess.check_output(["pdfinfo", str(pdf)], text=True)
    assert int(re.search(r"Pages:\s+(\d+)", info)[1]) == manifest["slide_count"]
    tree = ET.fromstring(
        subprocess.check_output(["pdftotext", "-bbox", str(pdf), "-"], text=True)
    )
    ns = {"x": "http://www.w3.org/1999/xhtml"}
    word_count = 0
    for i, page in enumerate(tree.findall(".//x:page", ns), 1):
        w, h = float(page.attrib["width"]), float(page.attrib["height"])
        for word in page.findall(".//x:word", ns):
            a = word.attrib
            assert 5 <= float(a["xMin"]) < float(a["xMax"]) <= w - 5, (i, word.text)
            assert 5 <= float(a["yMin"]) < float(a["yMax"]) <= h - 5, (i, word.text)
            word_count += 1
    fonts = subprocess.check_output(["pdffonts", str(pdf)], text=True).splitlines()[2:]
    assert fonts and all(
        re.search(r"\byes\s+(?:yes|no)\s+(?:yes|no)\s+\d+\s+\d+$", line)
        for line in fonts
    )
    text = subprocess.check_output(["pdftotext", "-layout", str(pdf), "-"], text=True)
    for token in (
        "VadCLIP",
        "SmolVLM",
        "SigLIP2",
        "MobileCLIP",
        "Linear probe",
        "ROC-AUC",
        "FPR",
        "Not run",
        "0.633",
        "0.976",
    ):
        assert token in text, token
    summary_pdf = SLIDES / "training_stages_summary_20261020.pdf"
    summary = Presentation(SLIDES / "training_stages_summary_20261020.pptx")
    assert len(summary.slides) == len(manifest["summary_layout"]) == 1
    planned = manifest["summary_layout"][0]
    summary_text = normalized(
        " ".join(
            shape.text for shape in summary.slides[0].shapes if shape.has_text_frame
        )
    )
    for item in planned["items"]:
        if item["kind"] == "text":
            assert normalized(item["value"]) in summary_text
            assert len(item["value"].splitlines()) * item["leading"] <= item["h"] + 0.01
    assert normalized(planned["notes"]) in normalized(
        summary.slides[0].notes_slide.notes_text_frame.text
    )
    summary_info = subprocess.check_output(["pdfinfo", str(summary_pdf)], text=True)
    assert int(re.search(r"Pages:\s+(\d+)", summary_info)[1]) == 1
    summary_tree = ET.fromstring(
        subprocess.check_output(
            ["pdftotext", "-bbox", str(summary_pdf), "-"], text=True
        )
    )
    for page in summary_tree.findall(".//x:page", ns):
        w, h = float(page.attrib["width"]), float(page.attrib["height"])
        for word in page.findall(".//x:word", ns):
            a = word.attrib
            assert 5 <= float(a["xMin"]) < float(a["xMax"]) <= w - 5, word.text
            assert 5 <= float(a["yMin"]) < float(a["yMax"]) <= h - 5, word.text
    if args.render:
        pages, sheets = render_contact_sheets(pdf, DETAIL / "previews")
        assert pages == manifest["slide_count"]
        (DETAIL / "previews/index.md").write_text(
            "# All rendered slides\n\n"
            + "\n\n".join(f"![Slide review]({s})" for s in sheets)
            + "\n"
        )
        subprocess.run(
            [
                "pdftoppm",
                "-scale-to",
                "1600",
                "-singlefile",
                "-png",
                str(summary_pdf),
                str(DETAIL / "summary"),
            ],
            check=True,
        )
    result = {
        "status": "passed",
        "slides": manifest["slide_count"],
        "metric_rows_compared_with_original_results": metric_count,
        "dataset_rows_recounted": dataset_count,
        "pptx_text_and_notes_checked": len(presentation.slides),
        "summary_slides_checked": 1,
        "pdf_words_within_bounds": word_count,
        "fonts_embedded": True,
        "source_files_checked": len(manifest["sources_sha256"]),
        "new_model_inference": False,
        "artifacts_sha256": {
            str(p.relative_to(ROOT)): digest(p)
            for p in (
                pdf,
                summary_pdf,
                SLIDES / "training_stages_summary_20261020.pptx",
                SLIDES / (STEM + ".pptx"),
                DETAIL / "metrics.csv",
                DETAIL / "dataset_counts.csv",
                Path(__file__),
            )
        },
    }
    (DETAIL / "validation.json").write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
