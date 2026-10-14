"""Build editable slides and a matching vector PDF from completed evaluations.

Install requirements.txt in a presentation environment, then run this file.
No model inference is performed. Existing evaluation artifacts are read only.
"""

from __future__ import annotations

import hashlib
import json
import math
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
RUN = ROOT / "results/gpu_all_models_20261013"
W, H = 960, 540
NAVY = "142C40"
TEAL = "007E80"
ORANGE = "C86535"
RED = "AB4048"
GRAY = "526471"
LINE = "DCE3E5"
WHITE = "FFFFFF"
PALE = "F0F6F6"
WARM = "F8F2EA"
BG = "FAFBF9"
NAMES = ["mobileclip2-s2", "siglip2-b16-224", "smolvlm-500m", "minicpm-v4"]
SHORT = ["MobileCLIP2-S2", "SigLIP2-B/16-224", "SmolVLM-500M", "MiniCPM-V 4.0"]
COLORS = [TEAL, "267AA6", ORANGE, "775791"]
SOURCES: dict[str, str] = {}
SLIDES: list[dict] = []


def read(path: Path):
    SOURCES[str(path.relative_to(ROOT))] = hashlib.sha256(path.read_bytes()).hexdigest()
    return json.loads(path.read_text())


def summary(cohort, split, model):
    return read(RUN / cohort / split / model / "summary.json")


pdfmetrics.registerFont(
    TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
)
pdfmetrics.registerFont(
    TTFont("DejaVu-Bold", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
)


def wrap(value, width, size, bold=False):
    font = "DejaVu-Bold" if bold else "DejaVu"
    lines = []
    for paragraph in value.split("\n"):
        if not paragraph:
            lines.append("")
            continue
        current = ""
        for word in paragraph.split():
            if pdfmetrics.stringWidth(word, font, size) > width:
                if current:
                    lines.append(current)
                    current = ""
                chunks = []
                part = ""
                for char in word:
                    if pdfmetrics.stringWidth(part + char, font, size) > width:
                        chunks.append(part)
                        part = char
                    else:
                        part += char
                lines.extend(chunks)
                current = part
                continue
            candidate = (current + " " + word).strip()
            if current and pdfmetrics.stringWidth(candidate, font, size) > width:
                lines.append(current)
                current = word
            else:
                current = candidate
        if current:
            lines.append(current)
    return lines


def box(slide, x, y, w, h, fill=WHITE, border=None):
    slide["items"].append(
        {"kind": "rect", "x": x, "y": y, "w": w, "h": h, "fill": fill, "border": border}
    )


def line(slide, x1, y1, x2, y2, color=LINE, width=1):
    slide["items"].append(
        {
            "kind": "line",
            "x": x1,
            "y": y1,
            "x2": x2,
            "y2": y2,
            "color": color,
            "width": width,
        }
    )


def text(slide, x, y, w, h, value, size=18, color=NAVY, bold=False, align="left"):
    lines = wrap(str(value), w, size, bold)
    leading = size * 1.23
    if len(lines) * leading > h + 0.01:
        raise ValueError(
            f"Slide {slide['title']!r}: text does not fit {w}×{h}: {value!r} "
            f"({len(lines)} lines at {size} pt)"
        )
    slide["items"].append(
        {
            "kind": "text",
            "x": x,
            "y": y,
            "w": w,
            "h": h,
            "value": "\n".join(lines),
            "size": size,
            "color": color,
            "bold": bold,
            "align": align,
            "leading": leading,
        }
    )


def footer(slide, index, source):
    line(slide, 42, 505, 918, 505)
    text(
        slide,
        42,
        514,
        835,
        14,
        f"CAPSTONE S25  /  EVALUATION 13 OCT 2026  •  {source}",
        8.5,
        GRAY,
    )
    text(slide, 882, 512, 36, 17, f"{index:02d}", 11, TEAL, True, "right")


def new(category, title, subtitle, source, notes):
    slide = {
        "title": title,
        "category": category,
        "notes": notes,
        "source": source,
        "items": [],
    }
    SLIDES.append(slide)
    box(slide, 0, 0, W, H, BG)
    box(slide, 42, 26, 28, 4, TEAL)
    text(slide, 80, 20, 838, 20, category.upper(), 11, TEAL, True)
    text(slide, 42, 58, 876, 80, title, 30, NAVY, True)
    text(slide, 42, 128, 876, 22, subtitle, 13.5, GRAY)
    footer(slide, len(SLIDES), source)
    return slide


def card(slide, x, y, w, h, title, body, color=TEAL, size=17):
    box(slide, x, y, w, h, WHITE, LINE)
    box(slide, x, y, 4, h, color)
    text(slide, x + 18, y + 14, w - 34, 50, title, 20, color, True)
    text(slide, x + 18, y + 70, w - 34, h - 80, body, size, NAVY)


def banner(slide, y, value, color=TEAL, size=16, h=44):
    box(slide, 42, y, 876, h, PALE if color == TEAL else WARM)
    box(slide, 42, y, 4, h, color)
    text(slide, 57, y + 9, 843, h - 12, value, size, color, True)


def stat(slide, x, y, w, h, value, label, detail="", color=TEAL):
    box(slide, x, y, w, h, WHITE, LINE)
    text(slide, x + 16, y + 10, w - 32, 55, value, 36, color, True)
    text(slide, x + 16, y + 65, w - 32, 28, label, 16, NAVY, True)
    if detail:
        text(slide, x + 16, y + 96, w - 32, h - 103, detail, 13, GRAY)


def table(slide, x, y, widths, headings, rows, row_h=45, font=16, header_h=38):
    total = sum(widths)
    box(slide, x, y, total, header_h, NAVY)
    xx = x
    for width, label in zip(widths, headings):
        text(slide, xx + 10, y + 9, width - 18, header_h - 10, label, 13, WHITE, True)
        xx += width
    for i, row in enumerate(rows):
        yy = y + header_h + i * row_h
        box(slide, x, yy, total, row_h, WHITE if i % 2 == 0 else PALE)
        xx = x
        for j, (width, value) in enumerate(zip(widths, row)):
            text(
                slide,
                xx + 10,
                yy + 9,
                width - 18,
                row_h - 11,
                value,
                font,
                NAVY,
                j == 0,
            )
            xx += width
        line(slide, x, yy + row_h, x + total, yy + row_h)


def p(value):
    return f"{value * 100:.2f}%"


def metric(value):
    return "—" if value is None else f"{value:.3f}"


def build_content():
    state = read(RUN / "evaluation.json")
    assert state["status"] == "complete_for_available_tasks"
    measured = 0
    for benchmark in state["benchmarks"]:
        for name in NAMES:
            result = summary(benchmark["cohort"], benchmark["split"], name)
            assert result["status"] == "ok" and result["errors"] == 0
            assert result["input_policy"]["device"] == "cuda"
            assert result["throttle_sleep_seconds"] == 0
            measured += result["completed_inferences"]
    assert measured == 13400
    smoke = {n: summary("smoking_frame", "dev", n) for n in NAMES}
    uncertainty = read(RUN / "smoking_frame/dev/improvements.json")
    latency = {
        n: read(ROOT / "results/cpu_latency_20261013" / n / "summary.json")
        for n in NAMES
    }
    latency_run = read(ROOT / "results/cpu_latency_20261013/run.json")
    vad = read(RUN / "vadclip/adapter/eval.json")
    vad_training = read(RUN / "vadclip/adapter/metrics.json")
    vad_baselines = read(RUN / "vadclip/baselines/baselines.json")
    sweep = read(ROOT / "results/vadclip_dev_sweep_20261013/selection.json")
    assert sum(s["completed_inferences"] for s in smoke.values()) == 13112
    assert len(latency_run["samples"]) == 8
    assert all(
        s["status"] == "ok" and s["throttle_sleep_seconds"] == 0 for s in smoke.values()
    )
    for path in [
        ROOT / "reports/workflow_review_20261013.md",
        ROOT / "reports/model_comparison_20261013.csv",
        ROOT / "reports/hardware_20261013.json",
    ]:
        SOURCES[str(path.relative_to(ROOT))] = hashlib.sha256(
            path.read_bytes()
        ).hexdigest()

    s = new(
        "Results briefing",
        "CLIP & VLM anomaly detection",
        "Findings, methodology, testing approach and evaluation",
        "Combined GPU comparison + workflow review",
        "This briefing uses the completed 13 October 2026 runs, not the older report open in the IDE. It summarizes MobileCLIP2-S2, SigLIP2-B/16-224, SmolVLM-500M-Instruct, MiniCPM-V 4.0 and the separate supervised VadCLIP smoking adaptation. No five-task deployment winner is supported. Sources: reports/workflow_review_20261013.md and results/gpu_all_models_20261013/report.md.",
    )
    box(s, 42, 168, 876, 109, NAVY)
    text(
        s,
        62,
        185,
        830,
        83,
        "Fast inference did not translate into\nreliable event detection.",
        30,
        WHITE,
        True,
    )
    for x, value, label in [
        (42, "4 + 1", "Image models + VadCLIP"),
        (340, "3,350", "Confirmed inputs / model"),
        (638, "13,400", "GPU quality predictions"),
    ]:
        stat(s, x, 298, 280, 105, value, label)
    text(
        s,
        44,
        427,
        872,
        54,
        "Five target events: smoking, littering, loitering, rough sleeping and unattended items. Only smoking and unattended items currently have usable confirmed cohorts.",
        17,
        GRAY,
    )

    s = new(
        "One-slide summary",
        "What we learned — and how we tested it",
        "Current-protocol findings; exploratory evidence, with no deployment winner",
        "Combined GPU report; CPU latency pilot; dataset audit",
        "Standalone summary slide. Methodology: source review, deduplication, recording-group splits, identical inputs and dev-only thresholds. GPU quality: four models × 3,350 inputs = 13,400 predictions, without pacing. VadCLIP has a different supervised offline split. CPU latency: same eight dev images, three repeats, four logical CPUs, float32, one worker, batch one. Smoking baseline recall: MobileCLIP 0.56%, SigLIP 0%, SmolVLM 5.41%, MiniCPM 0%. The three other tasks need verified event and normal labels; CAVIAR testing has only one positive source sequence per split.",
    )
    card(
        s,
        42,
        160,
        427,
        154,
        "Methodology",
        "Review event labels; deduplicate inputs; keep recording groups together; fit thresholds on development data.",
        TEAL,
        16,
    )
    card(
        s,
        491,
        160,
        427,
        154,
        "Testing approach",
        "GPU quality: 13,400 predictions, no pacing. Separate four-core CPU latency pilot. VadCLIP is a supervised study.",
        TEAL,
        16,
    )
    card(
        s,
        42,
        332,
        427,
        154,
        "Findings",
        "Smoking recall ≤5.41% for current image-model prompts. SmolVLM flags every unattended input; MiniCPM stays normal.",
        ORANGE,
        16,
    )
    card(
        s,
        491,
        332,
        427,
        154,
        "Evaluation limits",
        "Smoking is development-only. Unattended testing is tiny. Three tasks lack verified event labels and matched normals.",
        ORANGE,
        16,
    )

    s = new(
        "Research design",
        "Compare model families on shared inputs",
        "Four pretrained image models; VadCLIP uses a different supervised protocol",
        "anomaly_eval/models.py; combined run metadata",
        "CLIP models score image/text similarity, with cached native text features; MobileCLIP2-S2 is loaded through its OpenCLIP checkpoint and reparameterized for inference. SigLIP uses independent sigmoid scores and anomaly-minus-normal raw-logit margins. VLMs use native chat templates and deterministic JSON-label requests. MiniCPM uses the publisher's chat interface, one slice in this baseline. VadCLIP uses a frozen official CLIP-B/16 encoder plus a trained offline adapter. Publisher references: https://github.com/apple-aiml-research/ml-mobileclip ; https://huggingface.co/google/siglip2-base-patch16-224 ; https://huggingface.co/HuggingFaceTB/SmolVLM-500M-Instruct ; https://huggingface.co/openbmb/MiniCPM-V-4 ; https://github.com/nwpu-zxr/VadCLIP .",
    )
    card(
        s,
        42,
        166,
        280,
        244,
        "CLIP similarity",
        "MobileCLIP2-S2\nSigLIP2-B/16-224\n\nContinuous anomaly scores. Native image processing and cached text features.",
        TEAL,
        17,
    )
    card(
        s,
        340,
        166,
        280,
        244,
        "Generative VLMs",
        "SmolVLM-500M-Instruct\nMiniCPM-V 4.0\n\nNative chat formatting. Hard labels, raw responses and abstention tracking.",
        ORANGE,
        17,
    )
    card(
        s,
        638,
        166,
        280,
        244,
        "VadCLIP adaptation",
        "Frozen CLIP-B/16 features plus a trained temporal adapter.\n\nFrame supervision and future context; offline smoking study.",
        "775791",
        17,
    )
    banner(
        s,
        431,
        "Compare image models within a cohort. Keep the supervised VadCLIP results separate.",
        h=48,
    )

    s = new(
        "Dataset readiness",
        "Only two tasks have usable event labels",
        "Downloaded files do not automatically establish the target anomaly",
        "reports/data_readiness_20261013.json; dataset preparation",
        "Smoking video COCO: 3,278 confirmed frames, 16 known camera/day groups, 1,422 positives and 1,856 normals; development only. An additional 936 CCTV frames are all positive and need source genealogy and matched controls. CAVIAR supplies 20 dev and 16 test endpoints; window sheets share the same endpoints. There are three source groups and only one positive source sequence per split. Littering action boxes, loitering person boxes and rough-sleeping posture boxes are excluded from confirmed event scoring. Rough sleeping has 224 image files but only 112 unique byte/RGB hashes. Sources: reports/data_readiness_20261013.json, data/evaluation_all/dataset.json and the combined report.",
    )
    table(
        s,
        42,
        164,
        [143, 153, 342, 238],
        ["Target", "Local material", "Annotation meaning", "Scoring readiness"],
        [
            [
                "Smoking",
                "3,278 frames",
                "Smoking / Not Smoking; 16 groups",
                "Development only",
            ],
            [
                "Unattended items",
                "36 endpoints\n+ 36 windows",
                "Reviewed CAVIAR events / controls",
                "Small dev/test cohort",
            ],
            [
                "Littering",
                "6,721 images",
                "Action boxes; rubbish disposal unverified",
                "Needs event / normal review",
            ],
            [
                "Loitering",
                "1,675 images\n95 videos",
                "Person boxes; no duration-event mapping",
                "Needs event intervals",
            ],
            [
                "Rough sleeping",
                "224 files\n112 unique",
                "Laying / standing pose proxies",
                "Needs verified sleeping events",
            ],
        ],
        row_h=54,
        font=13.5,
    )
    text(
        s,
        45,
        478,
        870,
        20,
        "Missing tasks remain unscored. No proxy labels or missing annotations are converted into event negatives.",
        12.5,
        ORANGE,
        True,
    )

    s = new(
        "Methodology",
        "Protect the data before measuring quality",
        "A repeatable pipeline from source annotations to matched model comparisons",
        "data.py; datasets/prepare.py; analysis.py",
        "The workflow validates annotations and source identity, then deduplicates exact bytes, decoded EXIF-corrected RGB pixels and checksum-qualified source moments. Conflicting labels and source leakage fail validation. Historical benchmark groups are development only. Recording groups remain together across frame/window protocols. Models see the same ordered cohort. Continuous thresholds are chosen on development and frozen for test. Metrics include failures and abstentions; recording-group bootstrap is paired across models. Model revisions, processors, definitions, code hashes, input hashes and raw responses remain in run artifacts.",
    )
    steps = [
        (
            "01",
            "Review labels",
            "Verify events, normal controls and recording identity.",
        ),
        ("02", "Standardize", "Build task-specific manifests and frame/window inputs."),
        (
            "03",
            "Dedup & split",
            "Check pixels and source moments; hold groups together.",
        ),
        (
            "04",
            "Run models",
            "Identical ordered inputs, native processors, saved responses.",
        ),
        (
            "05",
            "Calibrate & score",
            "Choose dev thresholds; freeze test; bootstrap groups.",
        ),
    ]
    for i, (number, title, body) in enumerate(steps):
        x = 42 + i * 178
        box(s, x, 172, 164, 206, WHITE, LINE)
        text(s, x + 14, 187, 136, 42, number, 29, TEAL, True)
        text(s, x + 14, 239, 136, 55, title, 17, NAVY, True)
        text(s, x + 14, 303, 136, 70, body, 12.8, GRAY)
        if i < 4:
            text(s, x + 164, 254, 14, 21, "›", 17, TEAL, True, "center")
    banner(
        s,
        411,
        "Leakage checks + shared cohorts + recorded provenance make the comparison auditable.",
        h=56,
    )

    s = new(
        "Testing approach",
        "Accuracy and latency are separate tests",
        "Use GPU capacity for quality; freeze CPU resources only for the latency pilot",
        "GPU run metadata; cpu_latency_20261013/run.json",
        "GPU: RTX 5000 Ada Generation with 32,760 MiB VRAM; isolated CUDA Torch 2.8.0+cu128, no system driver changes. All 3,350 prepared confirmed inputs per model; zero per-group cap and inference limit. No duty-cycle/camera-rate sleeps, no default accuracy RAM cap. CLIP float32 batches of 32; VLM unquantized bf16 single-input calls. Shared letterbox max-side 768 and 64-token deterministic generation. CPU pilot: original Torch 2.14.1+cpu on Threadripper PRO 7975WX, affinity 0–3 and four threads, float32, one worker, batch one. Same eight unattended-item dev images, three measured repeats, three warmups, 24 observations/model. The host budget does not emulate a different processor.",
    )
    card(
        s,
        42,
        164,
        427,
        287,
        "GPU accuracy evaluation",
        "RTX 5000 Ada • 32 GB VRAM\nAll 3,350 confirmed inputs / model\nNo pacing sleeps or default RAM cap\nCLIP: float32, batch 32\nVLMs: bfloat16, one input / call\n768-pixel letterbox; 64-token budget",
        TEAL,
        17,
    )
    card(
        s,
        491,
        164,
        427,
        287,
        "CPU latency pilot",
        "Threadripper PRO 7975WX\nFour fixed logical CPUs + four threads\nFloat32, one worker, batch one\nSame eight dev images, three repeats\nThree warmups; 24 observations / model\nDecode → inference → parse included",
        ORANGE,
        17,
    )
    text(
        s,
        45,
        470,
        870,
        25,
        "GPU timings are observational. CPU timing is a small host-specific replay, excluding capture and queueing.",
        12.5,
        GRAY,
    )

    s = new(
        "Evaluation",
        "Measure misses and false alarms together",
        "Accuracy alone can reward a model that always returns normal",
        "metrics.py; calibration.json; improvements.json",
        "Recall is TP/(TP+FN); precision is TP/(TP+FP); F1 summarizes precision and recall; false-positive rate is FP divided by labeled normal inputs. Unsupported labels are not zero-scored. AP and AUROC use genuine continuous scores, with noninterpolated AP and tied-score handling. VLM categorical outputs do not receive invented AP/AUROC. All failures/abstentions remain in categorical quality denominators; prediction coverage is separate. All main group analyses use 2,000 recording-group bootstrap draws, paired across models. Few CAVIAR groups make uncertainty exploratory. Raw argmax and dev-calibrated operating points are separate.",
    )
    cards = [
        ("Recall", "How many true event inputs were detected?", TEAL),
        (
            "Precision & F1",
            "How reliable are alarms, and how well is recall balanced?",
            TEAL,
        ),
        (
            "False-positive rate",
            "What fraction of normal inputs were falsely flagged?",
            ORANGE,
        ),
        ("AP & AUROC", "Continuous score ranking only. No invented VLM AUC.", TEAL),
        (
            "Coverage",
            "Invalid outputs and abstentions remain visible and count in quality.",
            ORANGE,
        ),
        (
            "Source-group intervals",
            "2,000 paired bootstrap draws over recordings, not isolated frames.",
            TEAL,
        ),
    ]
    for i, (title, body, color) in enumerate(cards):
        card(
            s,
            42 + (i % 3) * 298,
            161 + (i // 3) * 165,
            280,
            149,
            title,
            body,
            color,
            15.5,
        )

    s = new(
        "Finding 1 / smoking",
        "Current prompts miss most smoking positives",
        "3,278 development frames • 1,422 positive • 1,856 normal • 16 known source groups",
        "smoking_frame/dev summaries and group bootstrap",
        "These are current-protocol development results, not unseen-camera tests. MobileCLIP TP=8 FP=4 FN=1,414 TN=1,852. SigLIP all normal. SmolVLM TP=77 FP=114 FN=1,345 TN=1,732, with 14 invalid responses (four positive and ten normal); those are included in quality. MiniCPM all normal. Positive prevalence is 0.433801, so constant-score AP is 0.434. MobileCLIP AP 0.424 and SigLIP AP 0.443 are near that reference. Paired AP/AUROC intervals do not distinguish the CLIP models robustly. SmolVLM-minus-MobileCLIP F1 CI crosses zero. These measure the tested prompts/preprocessing, not all native capabilities.",
    )
    rows = []
    for name, label in zip(NAMES, SHORT):
        b = smoke[name]["metrics"]["binary"]
        rows.append(
            [
                label,
                p(b["precision"]),
                p(b["recall"]),
                metric(b["f1"]),
                p(b["false_positive_rate"]),
                metric(b["average_precision"]),
                metric(b["roc_auc"]),
            ]
        )
    table(
        s,
        42,
        176,
        [226, 113, 112, 84, 114, 109, 118],
        ["Model", "Precision", "Recall", "F1", "FPR", "AP", "AUROC"],
        rows,
        row_h=47,
        font=16,
    )
    box(s, 42, 425, 427, 60, PALE)
    text(
        s,
        57,
        438,
        395,
        44,
        "MobileCLIP detects 8 positives; SmolVLM detects 77 out of 1,422.",
        16,
        TEAL,
        True,
    )
    box(s, 491, 425, 427, 60, WARM)
    text(
        s,
        506,
        438,
        395,
        44,
        "Always-normal accuracy is 56.6%. Its smoking recall is zero.",
        16,
        ORANGE,
        True,
    )

    s = new(
        "Finding 2 / unattended items",
        "Constant answers do not show event detection",
        "Per test protocol: 4 positive + 12 normal endpoints; only 3 source groups",
        "unattended_items frame/window test summaries",
        "Frame and past-window protocols share the same endpoints; they are not independent event samples. Every split has only one positive source sequence. Raw frame F1: MobileCLIP 0, SigLIP 0, SmolVLM 0.4, MiniCPM 0. Raw window F1: 0.25, 0, 0.4, 0 respectively. SmolVLM predicts unattended items on all dev/test frame/window inputs. Its 0.4 test F1 equals the always-anomaly baseline, with FPR 1. MiniCPM predicts normal throughout. MobileCLIP window argmax finds one of four positives, with three false positives. MobileCLIP max-dev-F1 frame threshold obtains test recall 1, precision 1/3, F1 0.5, FPR 2/3. Dev FPR-budget thresholds give zero test recall for both CLIPs.",
    )
    rows = []
    for name, label in zip(NAMES, SHORT):
        bf = summary("unattended_items_frame", "test", name)["metrics"]["binary"]
        bw = summary("unattended_items_window", "test", name)["metrics"]["binary"]
        rows.append(
            [
                label,
                metric(bf["f1"]),
                metric(bw["f1"]),
                p(bw["recall"]),
                p(bw["false_positive_rate"]),
            ]
        )
    table(
        s,
        42,
        169,
        [250, 151, 154, 153, 168],
        ["Model", "Frame F1", "Window F1", "Window recall", "Window FPR"],
        rows,
        row_h=45,
        font=16,
    )
    text(
        s,
        45,
        407,
        870,
        44,
        "SmolVLM: every input is an alarm. MiniCPM: every input is normal. Both patterns fail to distinguish events from controls.",
        17,
        ORANGE,
        True,
    )
    text(
        s,
        45,
        460,
        870,
        33,
        "Calibration: MobileCLIP frame recall reaches 100%, but FPR is 66.7%; low-FPR thresholds detect no test positives.",
        12.5,
        GRAY,
    )

    s = new(
        "Finding 3 / VadCLIP",
        "VadCLIP shows limited score discrimination",
        "A separate supervised offline smoking reanalysis; not a pooled zero-shot ranking",
        "vadclip/adapter/eval.json; training metrics; baselines",
        "Frozen CLIP-B/16 encoder and supervised temporal adapter with future context. Train 9 sessions/1,910 frames; development 4 sessions/312 frames; test 3 sessions/1,056 frames. Test has 560 positives and 496 normals. Selected branch dual_mean, epoch one, dev-fitted threshold 0.7381770357355234. Test TP290 FP177 FN270 TN319; precision 0.620985, recall 0.517857, F1 0.564752, FPR 0.356855, AP 0.633230, AUROC 0.587403. Visual and dual_max branches are all-positive with AUROC 0.5 and F1 0.693069. Linear probe FPR 0.977823. This is not the official weakly supervised UCF-Crime/XD-Violence task, and historical smoking corpus inspection means the split is a controlled reanalysis, not a fresh external test.",
    )
    for x, title, count, groups in [
        (42, "Train", vad_training["n_train_frames"], 9),
        (340, "Development", vad_training["n_valid_frames"], 4),
        (638, "Test", vad["n_test_frames"], 3),
    ]:
        box(s, x, 165, 280, 88, WHITE, LINE)
        text(s, x + 16, 178, 248, 30, title, 19, TEAL, True)
        text(s, x + 16, 216, 248, 26, f"{count:,} frames / {groups} sessions", 16, NAVY)
    b = vad["metrics"][vad["selected_branch"]]
    stat(s, 42, 279, 280, 115, metric(b["aupr"]), "Average precision", color=TEAL)
    stat(s, 340, 279, 280, 115, p(b["recall"]), "Test recall", color=TEAL)
    stat(
        s,
        638,
        279,
        280,
        115,
        p(b["false_positive_rate"]),
        "False-positive rate",
        color=ORANGE,
    )
    text(
        s,
        46,
        419,
        867,
        66,
        f"AUROC {metric(b['auroc'])}; F1 {metric(b['f1'])}. Visual / dual_max branches collapse to all-positive scores. Adapter-only timings exclude the encoder and cannot be compared with end-to-end VLM latency.",
        16,
        GRAY,
    )
    assert vad_baselines["split"] == "test"

    s = new(
        "Finding 4 / CPU resources",
        "CLIP is cheaper in the CPU latency pilot",
        "Same 8 development images × 3 repeats • float32 • 4 logical CPUs • batch one",
        "cpu_latency_20261013 summaries and run metadata",
        "The CPU latency pilot ran separately after quality jobs completed. Native decoding, preprocessing, model inference and parsing are included. Load and three warmups are excluded from the 24 measured timings per model; peak worker RSS includes load/warmup. MobileCLIP p95 0.056316s RSS993.316MiB; SigLIP 0.137628s RSS1248.602MiB; SmolVLM 7.563337s RSS3350.336MiB; MiniCPM 8.176413s RSS18850.168MiB. This is a small replay on Threadripper PRO 7975WX, not a hardware emulator or deployment queue/capture test. Output lengths and alternative image protocols can change costs. Model label abbreviations are used in the chart.",
    )
    for x, title in [
        (42, "p95 latency / seconds (log scale)"),
        (491, "Peak worker RSS / GiB"),
    ]:
        box(s, x, 163, 427, 288, WHITE, LINE)
        text(s, x + 16, 179, 397, 31, title, 16.5, NAVY, True)
    left_x, left_w, right_x, right_w = 161, 252, 614, 245
    for i, (name, label, color) in enumerate(
        zip(NAMES, ["MobileCLIP", "SigLIP", "SmolVLM", "MiniCPM"], COLORS)
    ):
        y = 229 + i * 48
        seconds = latency[name]["latency_p95_seconds"]
        gib = latency[name]["peak_rss_mb"] / 1024
        text(s, 57, y - 2, 98, 24, label, 14, GRAY)
        text(s, 506, y - 2, 101, 24, label, 14, GRAY)
        span = (math.log10(seconds) + 2) / 3 * left_w
        box(s, left_x, y, span, 19, color)
        text(s, left_x + span + 5, y - 2, 50, 25, f"{seconds:.3f}", 12, NAVY, True)
        box(s, right_x, y, gib / 20 * right_w, 19, color)
        text(
            s,
            right_x + gib / 20 * right_w + 5,
            y - 2,
            51,
            25,
            f"{gib:.2f}",
            12,
            NAVY,
            True,
        )
    for tick in [0.01, 0.1, 1, 10]:
        xx = left_x + (math.log10(tick) + 2) / 3 * left_w
        line(s, xx, 414, xx, 420)
        text(s, xx - 16, 424, 33, 18, str(tick), 10, GRAY, align="center")
    threshold_x = left_x + 2 / 3 * left_w
    line(s, threshold_x, 216, threshold_x, 414, GRAY, 0.6)
    text(s, threshold_x - 35, 207, 75, 16, "1 s budget", 9, GRAY, align="center")
    for tick in [0, 5, 10, 15, 20]:
        xx = right_x + tick / 20 * right_w
        line(s, xx, 414, xx, 420)
        text(s, xx - 16, 424, 33, 18, str(tick), 10, GRAY, align="center")
    text(
        s,
        45,
        467,
        870,
        28,
        "Both CLIPs met a 1-second inference budget here. The VLMs did not. These costs are host- and protocol-specific.",
        12.5,
        GRAY,
    )

    s = new(
        "Workflow improvements",
        "Remove limits that hid evaluation capacity",
        "Full GPU quality runs; explicit sampling and measurement semantics",
        "Code changes documented in workflow_review_20261013.md",
        "Changes include explicit runtime device/dtype resolution throughout adapters and VadCLIP; preserving GPU visibility instead of forced CUDA_VISIBLE_DEVICES=''; native CLIP batching and isolated concurrent workers; CUDA synchronization and memory peaks; removing default accuracy RAM cap; increasing timeout to four hours; removing the hidden per-recording preparation cap; recording available/retained sampling counts; indexing and invalidating download receipt lookups; AP and richer quality diagnostics; native one-input deployment timing. The old prepared cap produced 308 inputs versus 3,350 uncapped. Both environments passed 144 tests, plus lint/format/shell checks. No additional model inference is performed by this presentation generator.",
    )
    rows = [
        [
            "Execution",
            "Forced CPU / GPU hidden",
            "Explicit CPU / CUDA placement + dtype",
        ],
        [
            "Prepared cohort",
            "Hidden 8-per-label/recording cap",
            "3,350 / 3,350 confirmed inputs retained",
        ],
        [
            "Resources",
            "Restrictive defaults",
            "No pacing / default RAM cap; GPU batches",
        ],
        [
            "Data lookup",
            "Repeated 3 MB receipt parsing",
            "Indexed once; invalidates on changes",
        ],
        [
            "Reporting",
            "Incomplete quality / timing context",
            "AP, recall, FPR, coverage and batch kind",
        ],
    ]
    table(
        s,
        42,
        170,
        [160, 320, 396],
        ["Area", "Earlier limitation", "Updated workflow"],
        rows,
        row_h=45,
        font=14.5,
    )
    banner(
        s,
        448,
        "144 tests passed in each environment; all 20 GPU model/cohort summaries succeeded.",
        h=45,
        size=15,
    )

    s = new(
        "Conclusion & next evaluation",
        "No five-task deployment winner is supported",
        "Fix event coverage and test fresh recordings before making a model choice",
        "Workflow review: limitations and next evaluation",
        "No present configuration has demonstrated useful recall at an acceptable false-alert rate across five anomalies. Next steps are to review littering/loitering/rough-sleeping events and task-specific controls with camera/source genealogy; define operational duration and ownership rules; evaluate detected-person/object crops with detector failures included; investigate native VLM image slicing and video/temporal protocols on development only; freeze prompts/thresholds and test new recordings/cameras; report event detection delay and alerts/hour, then test latency on actual deployment hardware. Historical smoking and already-inspected CAVIAR test inputs are development evidence for subsequent protocol changes. Improvements are proposed experiments, not measured gains.",
    )
    cards = [
        (
            "1  Review event data",
            "Verified positives + task-specific normals.\n\nRecording identity, camera genealogy, duration and ownership rules.",
            TEAL,
        ),
        (
            "2  Improve the protocol",
            "Test crops and native temporal / image-slicing inputs on development data.\n\nInclude detector failures in quality.",
            TEAL,
        ),
        (
            "3  Freeze & retest",
            "Fresh source recordings; a fixed false-alert budget.\n\nMeasure event delay, alerts/hour and real-device latency.",
            ORANGE,
        ),
    ]
    for i, (title, body, color) in enumerate(cards):
        card(s, 42 + i * 298, 177, 280, 255, title, body, color, 16)
    text(
        s,
        45,
        455,
        870,
        37,
        "Current findings describe this baseline protocol. They do not exhaust the models' native capabilities.",
        15,
        GRAY,
        True,
    )

    s = new(
        "Appendix / development diagnostics",
        "Prompt sensitivity and uncertainty matter",
        "Development experiments identify hypotheses; they do not validate deployment gains",
        "smoking improvements.json; prompt experiment; VadCLIP LR sweep",
        "All smoking intervals are recording-group bootstraps over 16 known groups, 2,000 draws. MobileCLIP F1 CI[0.002442,0.027028]; SigLIP[0,0]; SmolVLM[0.006328,0.172909]; MiniCPM[0,0]. SmolVLM-minus-MobileCLIP F1 CI is [-0.012861,0.166235]. The short-definition CPU float32 CLIP experiment changed only text definitions and used all 3,278 smoking dev frames: MobileCLIP recall0.533052 FPR0.493534 F10.489664 AP0.460063; SigLIP remains all normal. VadCLIP CPU-cache development-only LR sweep: 1e-5 AP0.369294, 1e-4 AP0.393464, 1e-3 AP0.424815, selected 1e-3 epoch1 without selecting on test labels. Only selected candidate tested, AP0.631507. Native image preprocessing and numerical environment remain explicit.",
    )
    rows = []
    for name, label in zip(NAMES, SHORT):
        ci = uncertainty["models"][name]["ci95"]["binary_f1"]
        rows.append([label, f"{ci[0]:.3f}–{ci[1]:.3f}"])
    table(
        s,
        42,
        170,
        [245, 182],
        ["Smoking model", "F1 group CI95"],
        rows,
        row_h=43,
        font=15,
    )
    card(
        s,
        491,
        170,
        427,
        212,
        "Short-prompt diagnostic",
        "MobileCLIP recall: 53.3%\nBut false-positive rate: 49.4%\nAP: 0.460; SigLIP still all normal.\n\nA changed decision bias is not a validated gain.",
        ORANGE,
        16,
    )
    chosen = sweep["selected"]
    text(
        s,
        45,
        406,
        870,
        69,
        f"VadCLIP learning-rate sweep: dev AP 0.369 / 0.393 / 0.425 for 1e-5 / 1e-4 / 1e-3. Selection stayed at {chosen['lr']}, epoch {chosen['best_epoch']}. SmolVLM–MobileCLIP paired F1 interval includes zero (−0.013 to 0.166).",
        15,
        GRAY,
    )

    s = new(
        "Appendix / threshold evaluation",
        "Development thresholds expose the tradeoff",
        "Thresholds were frozen before test; only 16 normal development endpoints per protocol",
        "unattended_items_frame/window/test/calibration.json",
        "The 5% FPR budget is an empirical development constraint. With 16 normal development endpoints, FPR increments are 6.25%, so the constraint requires zero observed dev false positives; it is not a population guarantee. Rows are test outcomes under frozen thresholds. Frame/window cohorts share endpoint labels and have only three source groups per split, with one positive sequence. Both CLIP budget thresholds yield zero test recall. All max-dev-F1 and budget operating points are shown, not only favorable ones. VLM hard label outputs have no score threshold fitting.",
    )
    rows = []
    for protocol in ["frame", "window"]:
        calibration = read(RUN / f"unattended_items_{protocol}/test/calibration.json")
        for name, label in zip(NAMES[:2], ["MobileCLIP", "SigLIP"]):
            model = calibration["models"][name]
            for op, op_label in [
                ("dev_max_f1", "Max dev F1"),
                ("dev_fpr_budget", "Dev FPR budget"),
            ]:
                b = model["operating_points"][op]["all_samples"]["binary"]
                rows.append(
                    [
                        protocol,
                        label,
                        op_label,
                        metric(model["thresholds"][op]),
                        p(b["recall"]),
                        p(b["false_positive_rate"]),
                        metric(b["f1"]),
                    ]
                )
    table(
        s,
        42,
        169,
        [90, 138, 214, 104, 110, 110, 110],
        ["Input", "Model", "Dev selection", "Threshold", "Recall", "FPR", "F1"],
        rows,
        row_h=33,
        font=12.5,
    )
    text(
        s,
        45,
        478,
        870,
        17,
        "Small staged source-sequence holdout: no inference about operational alerts/hour or unseen-camera robustness.",
        11.5,
        GRAY,
    )

    s = new(
        "Evidence & reproducibility",
        "Every claim traces back to a saved run",
        "Use the latest combined run; the older IDE report predates these comparisons",
        "Sources are linked in speaker notes and source_manifest.json",
        "Primary evidence: results/gpu_all_models_20261013/report.md and evaluation.json; per-model summary.json and predictions.jsonl; per-cohort improvements.json and calibration.json. Resource pilot: results/cpu_latency_20261013/run.json and per-model summaries. Dataset audit: reports/data_readiness_20261013.json. Interpretation and limitations: reports/workflow_review_20261013.md. Numeric export: reports/model_comparison_20261013.csv. Supplementary experiments: results/clip_prompt_dev_20261013/short_descriptions and results/vadclip_dev_sweep_20261013/selection.json. Reproduction on this workspace uses scripts/evaluate_gpu.sh run --device cuda --per-group 0 --batch-size 32 --vadclip. It creates a new results directory. Future five-task readiness uses --require-all-tasks; currently incomplete. Source hashes and the exact slide data are captured in source_manifest.json. Rendering uses python-pptx and ReportLab; the PDF is built from the same layout specification, not a PowerPoint export.",
    )
    card(
        s,
        42,
        168,
        427,
        250,
        "Evidence files",
        "Combined GPU report + raw predictions\nCPU latency metadata + summaries\nData readiness audit\nWorkflow review + numeric CSV\n\nFull file paths and sources are in the notes and manifest.",
        TEAL,
        16,
    )
    card(
        s,
        491,
        168,
        427,
        250,
        "Reproduce & review",
        "Run the isolated GPU helper with:\nrun --device cuda --per-group 0\n--batch-size 32 --vadclip\n\nFreeze the protocol before a fresh test. Inspect false alarms and misses, not only aggregate accuracy.",
        ORANGE,
        16,
    )
    banner(
        s,
        442,
        "Evidence is complete for available tasks; verified five-task event coverage remains incomplete.",
        h=45,
        size=15,
    )
    assert len(SLIDES) == 16


def check_geometry(slides):
    for slide in slides:
        for item in slide["items"]:
            if item["kind"] == "line":
                assert 0 <= item["x"] <= W and 0 <= item["x2"] <= W
                assert 0 <= item["y"] <= H and 0 <= item["y2"] <= H
            else:
                assert item["x"] >= 0 and item["y"] >= 0
                assert item["x"] + item["w"] <= W + 0.1, (slide["title"], item)
                assert item["y"] + item["h"] <= H + 0.1, (slide["title"], item)


def render_pptx(slides, path):
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Pt(W), Pt(H)
    presentation.core_properties.title = (
        "CLIP & VLM anomaly detection: findings and evaluation"
    )
    presentation.core_properties.subject = (
        "GPU quality evaluation and separate CPU latency pilot"
    )
    presentation.core_properties.author = "Capstone S25"
    for spec in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for item in spec["items"]:
            k = item["kind"]
            if k == "rect":
                shape = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE, *[Pt(item[v]) for v in ["x", "y", "w", "h"]]
                )
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(item["fill"])
                if item["border"]:
                    shape.line.color.rgb = RGBColor.from_string(item["border"])
                    shape.line.width = Pt(0.7)
                else:
                    shape.line.fill.background()
            elif k == "line":
                from pptx.enum.shapes import MSO_CONNECTOR

                shape = slide.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT,
                    Pt(item["x"]),
                    Pt(item["y"]),
                    Pt(item["x2"]),
                    Pt(item["y2"]),
                )
                shape.line.color.rgb = RGBColor.from_string(item["color"])
                shape.line.width = Pt(item["width"])
            else:
                shape = slide.shapes.add_textbox(
                    *[Pt(item[v]) for v in ["x", "y", "w", "h"]]
                )
                frame = shape.text_frame
                frame.clear()
                frame.margin_left = frame.margin_right = frame.margin_top = (
                    frame.margin_bottom
                ) = 0
                frame.word_wrap = False
                frame.vertical_anchor = MSO_ANCHOR.TOP
                for i, row in enumerate(item["value"].split("\n")):
                    paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                    paragraph.text = row
                    paragraph.font.name = "DejaVu Sans"
                    paragraph.font.size = Pt(item["size"])
                    paragraph.font.bold = item["bold"]
                    paragraph.font.color.rgb = RGBColor.from_string(item["color"])
                    paragraph.line_spacing = Pt(item["leading"])
                    paragraph.space_before = paragraph.space_after = Pt(0)
                    paragraph.alignment = {
                        "left": PP_ALIGN.LEFT,
                        "center": PP_ALIGN.CENTER,
                        "right": PP_ALIGN.RIGHT,
                    }[item["align"]]
        slide.notes_slide.notes_text_frame.text = spec["notes"]
    presentation.save(path)


def render_pdf(slides, path):
    pdf = canvas.Canvas(str(path), pagesize=(W, H))
    pdf.setTitle("CLIP & VLM anomaly detection: findings and evaluation")
    pdf.setAuthor("Capstone S25")
    for spec in slides:
        for item in spec["items"]:
            k = item["kind"]
            if k == "rect":
                pdf.setFillColor(HexColor("#" + item["fill"]))
                pdf.setLineWidth(0.7)
                if item["border"]:
                    pdf.setStrokeColor(HexColor("#" + item["border"]))
                pdf.rect(
                    item["x"],
                    H - item["y"] - item["h"],
                    item["w"],
                    item["h"],
                    fill=1,
                    stroke=int(bool(item["border"])),
                )
            elif k == "line":
                pdf.setStrokeColor(HexColor("#" + item["color"]))
                pdf.setLineWidth(item["width"])
                pdf.line(item["x"], H - item["y"], item["x2"], H - item["y2"])
            else:
                font = "DejaVu-Bold" if item["bold"] else "DejaVu"
                pdf.setFont(font, item["size"])
                pdf.setFillColor(HexColor("#" + item["color"]))
                ascent = pdfmetrics.getAscent(font, item["size"])
                for i, row in enumerate(item["value"].split("\n")):
                    yy = H - item["y"] - ascent - i * item["leading"]
                    if item["align"] == "center":
                        pdf.drawCentredString(item["x"] + item["w"] / 2, yy, row)
                    elif item["align"] == "right":
                        pdf.drawRightString(item["x"] + item["w"], yy, row)
                    else:
                        pdf.drawString(item["x"], yy, row)
        pdf.showPage()
    pdf.save()


def main():
    build_content()
    check_geometry(SLIDES)
    standalone = deepcopy(SLIDES[1])
    for item in standalone["items"]:
        if item["kind"] == "text" and item["x"] == 882 and item["y"] == 512:
            item["value"] = "01"
    for stem, slides in [
        ("model_evaluation_20261013", SLIDES),
        ("summary_slide_20261013", [standalone]),
    ]:
        render_pptx(slides, OUT / (stem + ".pptx"))
        render_pdf(slides, OUT / (stem + ".pdf"))
    notes = [
        "# Presentation speaker notes",
        "",
        "Evaluation: 13 October 2026. Sources are relative to the repository root.",
        "",
    ]
    for index, slide in enumerate(SLIDES, 1):
        notes += [f"## {index:02d}. {slide['title']}", "", slide["notes"], ""]
    (OUT / "speaker_notes.md").write_text("\n".join(notes))
    (OUT / "source_manifest.json").write_text(
        json.dumps(
            {
                "evaluation_date": "2026-10-13",
                "slide_count": len(SLIDES),
                "sources_sha256": SOURCES,
                "build_script_sha256": hashlib.sha256(
                    Path(__file__).read_bytes()
                ).hexdigest(),
                "layout": SLIDES,
                "pdf_rendering": "Same measured layout as PowerPoint, rendered independently with embedded fonts.",
            },
            indent=2,
        )
        + "\n"
    )
    print(
        f"Created {len(SLIDES)} slides, editable PPTX + PDF, and a standalone summary slide."
    )


if __name__ == "__main__":
    main()
