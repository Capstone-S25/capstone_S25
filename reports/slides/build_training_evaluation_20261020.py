"""Build a training-only four-model deck from preserved experiment artifacts."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path

from pptx import Presentation
from reportlab import rl_config

from reports.slides import build_presentation as d

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
DETAIL = OUT / "training_evaluation_20261020"
STEM = "model_training_evaluation_20261020"
CLIP = ROOT / "results/clip_finetuning_20261015"
SMOL = ROOT / "results/smolvlm_finetuning_20261019"
STRATEGY = ROOT / "results/strategy_study_20261016"
PILOT = ROOT / "results/video_pilot_20261016"
LATEST = ROOT / "results/generalization_pilot_20261020"
VAD = ROOT / "results/gpu_all_models_20261013/vadclip"
MODELS = ("mobileclip2-s2", "siglip2-b16-224")
SHORT = {MODELS[0]: "MobileCLIP", MODELS[1]: "SigLIP2"}
TASKS = ("smoking", "littering", "loitering", "rough_sleeping", "unattended_items")
NAMES = dict(
    zip(
        TASKS,
        (
            "Smoking / CCTV",
            "Littering / IWDD",
            "Loitering",
            "Public sleep",
            "Unattended",
        ),
    )
)
SOURCES = dict(
    zip(
        TASKS,
        (
            "smoking_cctv",
            "mivia_iwdd",
            "street_scene",
            "commons_visible_sleeping",
            "unattended_reviewed",
        ),
    )
)
METHODS = ("zero_shot", "linear_probe", "finetuned")
METHOD_NAMES = dict(zip(METHODS, ("Zero-shot", "Linear probe", "Visual tune")))
POINTS = ("raw", "dev_max_f1", "dev_fpr_budget")
EXPORTS = []
DATASETS = []


def register(path):
    d.SOURCES[str(path.relative_to(ROOT))] = hashlib.sha256(
        path.read_bytes()
    ).hexdigest()


def load(path):
    return d.read(path)


def rows(path):
    register(path)
    return [json.loads(s) for s in path.read_text().splitlines()]


def p(value):
    return f"{100 * value:.1f}%"


def f(value):
    return f"{value:.3f}"


def readable(value):
    text = str(value)
    text = re.sub(
        r"\b(seed|epoch|rank|alpha|dropout|LR|batch|patience|length|window|confidence|microbatch|accumulation|decay|warmup|max)(?=\d)",
        r"\1 ",
        text,
        flags=re.IGNORECASE,
    )
    text = re.sub(
        r"(?<=\d)(?=(?:train|dev|test|normal|positive|examples|frames|views|epochs|images|clips|photos|parameters|iterations|steps|units)\b)",
        " ",
        text,
    )
    text = re.sub(r",(?=[A-Za-z])", ", ", text)
    text = re.sub(r";(?=\S)", "; ", text)
    return re.sub(r"%(?=[A-Za-z])", "% ", text)


def new(
    category,
    title,
    subtitle,
    notes,
    source="Saved training, selection and held-out results • 13–20 Oct 2026",
):
    slide = d.new(
        category, readable(title), readable(subtitle), source, readable(notes)
    )
    for item in slide["items"]:
        if item["kind"] == "text":
            item["value"] = item["value"].replace(
                "EVALUATION 13 OCT 2026", "TRAINING REVIEW 20 OCT 2026"
            )
    return slide


def banner(slide, value, warn=False):
    d.banner(slide, 435, readable(value), d.ORANGE if warn else d.TEAL, size=13.5, h=57)


def table(
    category,
    title,
    subtitle,
    headers,
    values,
    widths,
    takeaway,
    notes,
    *,
    warn=False,
    font=12.5,
    row_h=None,
):
    assert sum(widths) == 876, widths
    slide = new(category, title, subtitle, notes)
    row_h = row_h or min(44, 228 // max(len(values), 1))
    assert 166 + 36 + row_h * len(values) <= 435, title
    d.box(slide, 42, 166, 876, 36, d.NAVY)
    x = 42
    for width, label in zip(widths, headers):
        d.text(slide, x + 8, 171, width - 16, 30, label, 12, d.WHITE, True)
        x += width
    for i, row in enumerate(values):
        y = 202 + i * row_h
        d.box(slide, 42, y, 876, row_h, d.WHITE if i % 2 == 0 else d.PALE)
        x = 42
        for j, (width, value) in enumerate(zip(widths, row)):
            d.text(
                slide,
                x + 8,
                y + 3,
                width - 16,
                row_h - 6,
                readable(value),
                font,
                d.NAVY,
                j == 0,
            )
            x += width
        d.line(slide, 42, y + row_h, 918, y + row_h)
    banner(slide, takeaway, warn)
    return slide


def paired(category, title, subtitle, left, right, takeaway, notes, *, warn=False):
    slide = new(category, title, subtitle, notes)
    for x, (heading, body), color in ((42, left, d.TEAL), (491, right, d.ORANGE)):
        d.card(
            slide,
            x,
            166,
            427,
            249,
            readable(heading),
            readable(body),
            color=color,
            size=16,
        )
    banner(slide, takeaway, warn)
    return slide


def export_metric(
    study, model, stage, task, source, point, metric, *, seed=42, threshold=None
):
    m = dict(metric)
    if "n_pos" in m:
        m["positive"], m["negative"] = m.pop("n_pos"), m.pop("n_neg")
        m["ap"], m["fpr"] = m.pop("aupr"), m.pop("false_positive_rate")
    m["accuracy"] = (m["tp"] + m["tn"]) / m["n"]
    m["balanced_accuracy"] = (m["recall"] + 1 - m["fpr"]) / 2
    EXPORTS.append(
        {
            "study": study,
            "model": model,
            "seed": seed,
            "stage": stage,
            "task": task,
            "source": source,
            "operating_point": point,
            "threshold": threshold,
            **{
                k: m[k]
                for k in (
                    "n",
                    "positive",
                    "negative",
                    "tp",
                    "fp",
                    "fn",
                    "tn",
                    "ap",
                    "auroc",
                    "precision",
                    "recall",
                    "f1",
                    "fpr",
                    "accuracy",
                    "balanced_accuracy",
                )
            },
        }
    )
    return m


class Evidence:
    def __init__(self):
        self.clip = {}
        self.selection = {}
        self.params = {}
        self.probes = {}
        for model in MODELS:
            for seed in (42, 43, 44):
                base = CLIP / f"seed_{seed}" / model
                self.clip[model, seed] = load(base / "float32_evaluation/results.json")
                self.selection[model, seed] = load(
                    base / "float32_evaluation/selection.json"
                )
                if seed == 42:
                    self.params[model] = load(base / "training_parameters.json")
                    self.probes[model] = load(base / "linear_probe_dev.json")
                    load(base / "state.json")
        self.smol_selection = {
            seed: load(SMOL / f"seed_{seed}/selection.json") for seed in (42, 43)
        }
        self.smol = {}
        self.smol_transfer = {}
        for route in (
            "original",
            "seed_42/lr_3e-05",
            "seed_42/lr_0.0001",
            "seed_43/lr_0.0001",
        ):
            name = route.replace("/", "__")
            self.smol[route] = load(SMOL / "evaluation" / name / "results.json")
            self.smol_transfer[route] = load(
                SMOL / "additional_transfer/evaluation" / name / "results.json"
            )
        self.smol_protocol = load(SMOL / "protocol.json")
        self.smol_histories = {
            route: load(SMOL / route / "history.json")
            for route in self.smol
            if route != "original"
        }
        self.strategy = {
            m: load(STRATEGY / m / "transfer_evaluation/per_source_results.json")
            for m in MODELS
        }
        self.strategy_routes = {
            m: load(STRATEGY / m / "selection.json") for m in MODELS
        }
        self.vad_train = load(VAD / "adapter/metrics.json")
        self.vad_eval = load(VAD / "adapter/eval.json")
        self.vad_base = load(VAD / "baselines/baselines.json")
        self.vad_manifest = load(
            ROOT / "data/models/vadclip_features_cuda/manifest.json"
        )
        self.vad_source = load(VAD / "source_validation.json")
        self.clip_rows = rows(ROOT / "data/clip_training_20261015/samples.jsonl")
        self.smol_rows = rows(SMOL / "samples.jsonl")
        self.extra_rows = rows(SMOL / "additional_transfer/samples.jsonl")
        self.new_rows = rows(
            ROOT / "data/generalization_pilot_20261020/xrf55/samples.jsonl"
        )
        self.dataset = load(ROOT / "data/clip_training_20261015/dataset.json")
        self.provenance = load(SMOL / "clip_comparison_provenance.json")
        self.latest = {m: load(LATEST / "smoking" / m / "results.json") for m in MODELS}
        self.latest_selections = {
            m: load(LATEST / "smoking" / m / "selection.json") for m in MODELS
        }
        self.latest_smol_calibration = {
            m: load(LATEST / "smolvlm_transfer" / m / "calibration.json")
            for m in ("original", "lora")
        }
        self.latest_smol = {
            m: load(LATEST / "smolvlm_transfer" / m / "results.json")
            for m in ("original", "lora")
        }
        load(LATEST / "smoking/protocol.json")
        load(LATEST / "smolvlm_transfer/protocol.json")
        for path in (
            CLIP / "verification.json",
            SMOL / "verification.json",
            LATEST / "verification.json",
        ):
            load(path)
        register(ROOT / "anomaly_eval/training/clip.py")
        register(ROOT / "anomaly_eval/vadclip/baselines.py")
        register(ROOT / "anomaly_eval/vadclip/train.py")

    def cm(self, model, method, task, point="dev_max_f1", seed=42, source=None):
        split = "external" if source == "stanford40" else "test"
        cell = self.clip[model, seed][method][split][task]
        return (
            cell["sources"][source][point]
            if source
            else cell["operating_points"][point]
        )

    def sm(self, task, source, method="routed", point="dev_max_f1", seed=42):
        sel = self.smol_selection[seed]
        route = (
            "original"
            if method == "original"
            else (
                sel["joint_adapter"] if method == "joint" else sel["task_routes"][task]
            )
        )
        if source in (
            "hmdb51_action_transfer",
            "commons_new_author_transfer",
            "meva_abandon_action_transfer",
        ):
            return self.smol_transfer[route][task]["sources"][source][point]
        split = "external" if source == "stanford40" else "test"
        return self.smol[route][split][task]["sources"][source][point]

    def st(self, model, task, source, point="dev_max_f1"):
        key = next(
            k for k in self.strategy[model] if k.endswith("/" + task + "/" + source)
        )
        return self.strategy[model][key]["selected"][point]


def dataset_counts(records, label):
    values = []
    for task in TASKS:
        cells = []
        for split in ("train", "dev", "test", "external"):
            group = [r for r in records if r["task"] == task and r["split"] == split]
            n, pos = len(group), sum(r["target"] for r in group)
            DATASETS.append(
                {
                    "study": label,
                    "task": task,
                    "split": split,
                    "n": n,
                    "positive": pos,
                    "negative": n - pos,
                }
            )
            cells.append(f"{n:,} / {pos:,}" if n else "—")
        values.append(
            ["Smoking (all sources)" if task == "smoking" else NAMES[task], *cells]
        )
    return values


def metric_values(label, stage, m):
    return [
        label,
        stage,
        f(m["ap"]),
        f(m["auroc"]),
        f(m["f1"]),
        p(m["precision"]),
        p(m["recall"]),
        p(m["fpr"]),
    ]


def metric_table(title, subtitle, values, takeaway, notes, **kwargs):
    return table(
        "Evaluation / stage metrics",
        title,
        subtitle,
        [
            "Model / source",
            "Stage",
            "AP ↑",
            "ROC-AUC ↑",
            "F1 ↑",
            "Prec. ↑",
            "Recall ↑",
            "FPR ↓",
        ],
        values,
        [174, 124, 84, 105, 80, 103, 103, 103],
        takeaway,
        notes,
        font=11.5,
        **kwargs,
    )


def stage_chart(e, model):
    slide = new(
        "Evaluation / matched training stages",
        SHORT[model] + ": every training stage on the same holdout",
        "15–16 Oct study • primary seed 42 • AP by target; selected routes are reported separately",
        "Zero-shot text uses the endpoint visual embedding; the probe and visual-tuned classifier use normalized last/mean/max/change features. Thus the first step adds supervision and representation pooling. Visual tuning versus the probe uses the same representation and dataset. These bars are all candidates, not a test-best deployment selection. AP does not determine a working cutoff. Values come from the final float32 evaluation, not the retained AMP diagnostic.",
    )
    colors = (d.GRAY, d.TEAL, d.ORANGE)
    for x, name, color in zip(
        (205, 420, 660), ("Zero-shot", "Linear probe", "Visual tune"), colors
    ):
        d.box(slide, x, 160, 14, 10, color)
        d.text(slide, x + 20, 155, 180, 25, name, 13)
    x0, width = 225, 590
    for j, task in enumerate(TASKS):
        y = 194 + j * 44
        d.text(slide, 42, y + 4, 175, 34, NAMES[task], 13, bold=True)
        for k, (method, color) in enumerate(zip(METHODS, colors)):
            value = e.cm(model, method, task)["ap"]
            d.box(slide, x0, y + k * 10, width * value, 7, color)
            d.text(slide, x0 + width * value + 6, y + k * 10 - 3, 70, 13, f(value), 9)
    banner(
        slide,
        "Training improves several rankings; smoking transfer and loitering remain inconsistent.",
        True,
    )


def scope_and_data(e):
    paired(
        "Training / scope",
        "Four models, every completed learning stage",
        "MobileCLIP2-S2 • SigLIP2-B/16-224 • SmolVLM-500M-Instruct • VadCLIP",
        (
            "What is compared",
            "Zero-shot / original model\nFrozen-feature linear probes\nPartial visual fine-tuning\nLanguage LoRA and temporal adapters\nLater region / source-data experiments",
        ),
        (
            "How to read the results",
            "Five target-specific binary tasks\nDataset and positive counts shown\nRanking + fixed-cutoff metrics\nPrimary development-selected routes\nUnavailable stages explicitly marked",
        ),
        "This deck covers training and accuracy evaluation; all numbers come from preserved runs.",
        "This is a dedicated training/evaluation presentation rather than the full project-history deck. The requested four families are included; MiniCPM is outside this deck. Six CLIP visual runs, three SmolVLM LoRA runs and the corrected supervised VadCLIP study are documented. Later classifier, region and source-data studies are separated from genuine changes to pretrained backbone weights. No model is retrained for this documentation build and no speed or latency measurements are included.",
    )
    table(
        "Training / stage coverage",
        "The stages differ by model architecture",
        "A stage is reported only when a corresponding saved experiment exists",
        ["Model", "Baseline", "Linear probe", "Weight adaptation", "Later stages"],
        [
            [
                "MobileCLIP",
                "Text margin",
                "Yes",
                "Last visual stage",
                "Crops / GRU / source fits",
            ],
            [
                "SigLIP2",
                "Text margin",
                "Yes",
                "Last visual layers",
                "Geometry / ensembles",
            ],
            [
                "SmolVLM",
                "Original binary query",
                "Not run",
                "Language-layer LoRA",
                "Joint / routed / transfer",
            ],
            [
                "VadCLIP",
                "CLIP-B/16 text",
                "Yes",
                "Temporal adapter",
                "Branch selection",
            ],
        ],
        [155, 180, 145, 190, 206],
        "No study unfreezes every parameter: visual tuning, LoRA and adapter training have distinct meanings.",
        "MobileCLIP updates 18,106,992 visual parameters plus 10,245 head parameters; SigLIP updates 14,176,512 plus 15,365. Earlier visual layers and text encoders remain frozen. SmolVLM updates 4,341,760 language projection parameters while its vision tower and connector remain frozen. VadCLIP trains about 12.642 million adapter parameters over frozen CLIP-B/16 features. A missing SmolVLM linear probe is not assigned another model's result.",
        font=12,
        row_h=47,
    )
    stage_values = []
    for model in MODELS:
        stage_values.append(
            [
                SHORT[model],
                "A: 646 / 393",
                *(f(e.cm(model, method, "smoking")["ap"]) for method in METHODS),
            ]
        )
    stage_values.append(
        [
            "SmolVLM",
            "A: 646 / 393",
            f(e.sm("smoking", "smoking_cctv", "original")["ap"]),
            "Not run",
            f(e.sm("smoking", "smoking_cctv", "joint")["ap"]),
        ]
    )
    stage_values.append(
        [
            "VadCLIP",
            "B: 1,056 / 560",
            f(e.vad_base["metrics"]["zeroshot"]["aupr"]),
            f(e.vad_base["metrics"]["linearprobe"]["aupr"]),
            f(e.vad_eval["metrics"]["dual_mean"]["aupr"]),
        ]
    )
    table(
        "Evaluation / four-model stage overview",
        "Smoking: all four model families, stage by stage",
        "AP overview • A and B are different tests • the adaptation type differs by architecture",
        [
            "Model",
            "Test N / positive",
            "Zero-shot / original",
            "Linear probe",
            "Weight-adapted",
        ],
        stage_values,
        [165, 190, 190, 156, 175],
        "A: camera-c1 holdout. B: session holdout across cameras. These rows do not establish a four-model winner.",
        "MobileCLIP/SigLIP stages are the original CLIP visual study: 3,258 train,1,633dev. SmolVLM original and joint LoRA use the expanded 3,973train/1,812dev corpus and matched binary prompting. They share the646-case test with the earlier CLIP stages but not the same training data. VadCLIP stages share their own1,056-frame test and9/4/3session split. Its temporal adapter, Smol language LoRA and CLIP visual fine-tunes are distinct forms of weight adaptation. SmolVLM has no executed linear-probe stage. This AP table must not be read as a matched four-model leaderboard.",
        warn=True,
        font=12,
        row_h=47,
    )
    table(
        "Data / study ledger",
        "Overall data coverage, without double counting studies",
        "Rows are study manifests and reuse sources • counts are examples, not independent events",
        ["Study", "Train", "Dev", "Test / transfer", "Total / unit"],
        [
            [
                "CLIP visual study",
                "3,258",
                "1,633",
                "1,801 + 806 external",
                "7,498 mixed examples",
            ],
            [
                "Strategy / Smol main",
                "3,973",
                "1,812",
                "2,025 + 806 external",
                "8,616 mixed examples",
            ],
            [
                "Additional transfer",
                "0",
                "0",
                "906 + 28 + 27",
                "961 clips/photos/windows",
            ],
            ["VadCLIP smoking", "1,910", "312", "1,056", "3,278 original frames"],
            ["HMDB adaptation", "530", "161", "205", "896 clips after purge"],
            ["20 Oct XRF55 proxy", "340", "160", "240", "740 scripted clips"],
        ],
        [240, 100, 100, 220, 216],
        "The 8,616-example corpus contains the earlier 7,498; these totals must not be added as unique data.",
        "The expansion adds 1,118 smoking photos: 715 training, 179 development and 224 test. The main SmolVLM corpus therefore has 3,973 training, 1,812 development, 2,025 test and 806 external examples. Additional transfer is wholly held out from that training. HMDB later changes from a 906-clip transfer cohort to grouped adaptation folds; ten title-part aliases are purged, leaving 896. The 3,278 VadCLIP frames reuse the original CCTV corpus under a different split. Views, repeated temporal endpoints and copied dataset exports do not create independent events.",
        font=11.5,
        row_h=37,
    )
    table(
        "Data / sources and meaning",
        "What the five targets are actually trained on",
        "Input units and annotation scope determine what each evaluation supports",
        ["Objective", "Sources used", "Main evaluation unit / limitation"],
        [
            [
                "Smoking",
                "CCTV; Stanford40; Kaggle; HMDB; XRF55",
                "Frames/photos/actions; XRF55 is a gesture proxy",
            ],
            [
                "Littering",
                "MIVIA-IWDD",
                "Whole waste-dumping clips; broader than small littering",
            ],
            [
                "Loitering",
                "Street Scene track/type labels",
                "20-second windows; one camera, supervised re-split",
            ],
            [
                "Public sleep",
                "Commons photos and matched controls",
                "Apparent visible sleep; provisional review",
            ],
            [
                "Unattended",
                "CAVIAR / ABODA; MEVA transfer",
                "Overlapping endpoints; ownership/state labels limited",
            ],
        ],
        [150, 325, 401],
        "Primary-action labels, person boxes and posture labels are not interchangeable with verified events.",
        "Source counts in the primary CLIP manifest: CCTV smoking 3,278; Stanford40 1,406; IWDD 435; Street Scene 496; apparent public sleep 113; unattended endpoints 1,770. Subsequent photo expansion adds 1,118. IWDD acquired 498/500 videos; two missing training videos and 63 training videos linked to official test sources are excluded, leaving 274/61/100. Old person-only loitering exports, duplicate lying/standing images and foreground-mask datasets are not treated as complete target supervision. Public-sleep labels describe appearance/context, with no inference about housing status.",
        font=12,
        row_h=43,
    )
    table(
        "Data / original CLIP learning corpus",
        "Primary CLIP dataset: total / positive by target",
        "7,498 target-specific examples • 12,478 distinct view images • one known objective per row",
        ["Target", "Train N / +", "Dev N / +", "Test N / +", "External N / +"],
        dataset_counts(e.clip_rows, "clip_visual_15oct"),
        [240, 159, 159, 159, 159],
        "Only the annotated objective contributes loss; the other four labels remain unknown.",
        "Smoking trains on camera c4 and Stanford40 official-training photos, develops on c2 and tests c1; the external 806 photos preserve Stanford40 official-test membership. IWDD retains its 100 official test clips. Street Scene recordings stay within a split, but all originate from one camera. Sleep splits use author/perceptual groups. Unattended training has two source groups, development one and historical test two camera families. Static targets use one endpoint, temporal tasks four chronological views, and littering eight uniform clip views.",
        row_h=42,
    )
    table(
        "Data / expanded learning corpus",
        "Expanded strategy and SmolVLM dataset: total / positive",
        "8,616 examples • 13,596 distinct decoded images • same expanded train/dev splits",
        ["Target", "Train N / +", "Dev N / +", "Test N / +", "External N / +"],
        dataset_counts(e.smol_rows, "expanded_strategy_smol_19oct"),
        [240, 159, 159, 159, 159],
        "The main expansion is smoking photos: +715 train, +179 dev and +224 held-out photos.",
        "The original CLIP visual fine-tuning study and this expanded corpus do not have identical training data. The 16 October CLIP strategy routes and 19 October SmolVLM study do share these expanded training/development splits and exactly matched held-out IDs/labels. That is the basis for the later three-model pipeline comparison. Main held-out cases total 2,831; additional transfer adds 961, giving 3,792. Neither count equals independent people, cameras or events.",
        row_h=42,
    )
    vad_rows = []
    for split in ("train", "dev", "test"):
        sessions = [e.vad_source["sessions"][key] for key in e.vad_manifest[split]]
        n = sum(s["n_frames"] for s in sessions)
        pos = sum(s["n_positive"] for s in sessions)
        vad_rows.append([split, len(sessions), n, pos, n - pos, p(pos / n)])
        DATASETS.append(
            {
                "study": "vadclip_sessions_13oct",
                "task": "smoking",
                "split": split,
                "n": n,
                "positive": pos,
                "negative": n - pos,
            }
        )
    table(
        "Data / VadCLIP split",
        "VadCLIP uses a separate session-held-out smoking split",
        "Frozen CLIP-B/16 features • original 3,278 annotated frames across 16 camera/date sessions",
        ["Split", "Sessions", "Frames", "Smoking +", "Normal −", "Prevalence"],
        vad_rows,
        [160, 120, 150, 150, 150, 146],
        "Train and test share physical cameras across dates; this is not the CLIP camera-c1 holdout.",
        "The split is 9 training, 4 development and 3 test sessions, grouped by camera/date. The 1,056 test frames contain 560 positives and 496 normals. Train and test both include camera IDs c1, c2 and c4 on different dates, so this is session generalization rather than unseen-camera generalization. The separate camera-c1 holdout used by the five-target CLIP/Smol studies has 646 frames and 393 positives. No arithmetic comparison between those scores is a matched four-model benchmark.",
        row_h=47,
    )
    table(
        "Data / external evidence",
        "Transfer cohorts answer different generalization questions",
        "These are separately reported cohorts; their labels and train/test roles are explicit",
        ["Cohort", "Held-out N / +", "Role", "What is still limited"],
        [
            [
                "Stanford40",
                "806 / 140",
                "Official photo test",
                "Photo domain contributes training",
            ],
            [
                "HMDB action transfer",
                "906 / 109",
                "Unseen-source test in strategy/Smol",
                "Other actions are non-exhaustive negatives",
            ],
            [
                "New-author public sleep",
                "28 / 17",
                "No training or selection",
                "25 author groups; provisional labels",
            ],
            [
                "MEVA abandonment",
                "27 / 15",
                "No training or selection",
                "Action windows, not persistent item state",
            ],
            [
                "XRF55 scene 4",
                "240 / 60",
                "20 Oct new-source proxy test",
                "One scene, three subject aliases; masked gestures",
            ],
        ],
        [225, 140, 230, 281],
        "The later HMDB 205-clip adaptation test and the earlier 906-clip transfer test must stay separate.",
        "HMDB transfer uses eight uniform views across full clips and includes 109 smoking plus 797 primary-action controls. The later adapted pilot uses sixteen views and 530/161/205 grouped clips after title-part alias purging. MEVA has 15 abandonment windows plus 12 pre-event controls across ten synchronized sessions/seven cameras, using source end annotations to define offline windows. It does not supply fully reviewed unattended-owner state. XRF55 train/dev/test aliases and scene blocks are disjoint; training inspection did not verify a visible cigarette, so it is weak gesture evidence.",
        font=11.5,
        row_h=43,
    )
    paired(
        "Evaluation / split and selection discipline",
        "Train, select and test play different roles",
        "Development data determine checkpoints, routes and cutoffs; test data report the frozen result",
        (
            "Training and validation",
            "Task / class / source balancing\nGrouped camera / recording / author splits\nUnknown targets have zero loss\nFile and decoded-pixel checks\nPrimary seeds specified in advance",
        ),
        (
            "Held-out reporting",
            "All stage results remain visible\nNo best-test-seed substitution\nNo test-fitted decision threshold\nHistorical results labelled retrospective\nPer-source denominators retained",
        ),
        "Seeds estimate optimization variation; they do not add cameras or independent event evidence.",
        "Historical CCTV, unattended intervals and many transfer cohorts were inspected in earlier rounds; those tests remain retrospective. Exact pixel disjointness does not establish complete semantic/source independence. Six CLIP visual runs use seeds 42/43/44; Smol uses two primary learning rates at seed42 plus one repeat at seed43. More source views, repeated photos or adjacent endpoints do not increase the count of independent target events. The XRF55 split was frozen before new source predictions. No latency experiment is included.",
        warn=True,
    )
    table(
        "Evaluation / metrics and operating points",
        "Use ranking and decision metrics together",
        "All metrics are binary and source-specific • accuracy is exported with balanced accuracy",
        ["Metric / point", "Meaning", "How it is used here"],
        [
            [
                "AP / PR area",
                "Precision–recall ranking across score cutoffs",
                "Higher is better; compare the same labelled cohort",
            ],
            [
                "ROC-AUC",
                "Pairwise positive/negative ranking",
                "Higher is better; show with AP and prevalence",
            ],
            [
                "Precision / recall / F1",
                "TP/(TP+FP); TP/P; harmonic mean",
                "Use a named, frozen cutoff; report confusion counts",
            ],
            [
                "FPR / accuracy",
                "FP/N; (TP+TN)/all",
                "FPR matters for normal scenes; accuracy can hide imbalance",
            ],
            [
                "Raw / dev-max-F1",
                "Zero margin / best development F1",
                "No test recalibration; main stage tables use dev-max-F1",
            ],
            [
                "Dev FPR budget",
                "At most 5% empirical development FPR",
                "Test FPR can exceed 5%; not an alert-per-hour rate",
            ],
        ],
        [210, 310, 356],
        "AP can rise while FPR becomes unacceptable. A classifier score is not an event probability.",
        "Main detailed stage tables use each method's maximum-development-F1 cutoff, frozen before held-out scoring. Supplementary tables explicitly use raw zero or the development 5% FPR point. Older CLIP and Smol cutoffs use pooled target development negatives; later 17/20 October smoking experiments constrain every development source. Accuracy and balanced accuracy are derived from saved TP/FP/FN/TN. All source thresholds and unrounded metrics are included in metrics.csv. AP/AUPR labels follow the stored implementations; VadCLIP AP belongs to its separate frame cohort. Endpoint FPR is not continuous false alerts per camera-hour.",
        font=11.5,
        row_h=38,
    )


def clip_training_and_stages(e):
    table(
        "Training / MobileCLIP and SigLIP2",
        "Linear probing and visual tuning change different weights",
        "Same primary dataset and development-selection rule • five binary objectives",
        ["Setting", "Frozen-feature probe", "Visual fine-tune"],
        [
            [
                "Trainable part",
                "Five logistic-regression classifiers",
                "Last visual stage/layers plus five heads",
            ],
            [
                "Representation",
                "Normalized last / mean / max / change",
                "Same pooled representation; learned visual features",
            ],
            [
                "Optimization",
                "C = 1, 10, 100; at most 2,000 LBFGS steps",
                "AdamW; visual LR 1e-5 or 3e-5; head LR 3e-4",
            ],
            [
                "Regularization",
                "Task/class/source/group sample weights",
                "Weight decay .01; dropout .1; clipping 1.0",
            ],
            [
                "Preserve pretraining",
                "Entire visual encoder frozen",
                "Feature MSE ×.1; parameter anchor ×.1",
            ],
            [
                "Budget / inputs",
                "C selected on dev macro source AP",
                "Batch 32; max 10 epochs; patience 3; BF16 train",
            ],
        ],
        [170, 333, 373],
        "Earlier visual layers, text encoders and batch-normalization statistics remain frozen.",
        "The probe pools ordered input image embeddings into normalized endpoint/mean/max/last-minus-first features. C is selected on development AP: primary MobileCLIP C10, SigLIP C100. Only one annotated target contributes BCE to visual tuning. Fine-tuning starts from the selected probe head, updates the recorded last-stage visual parameters, and applies feature distillation plus weight anchoring at coefficient0.1 each. Scene-preserving augmentation uses horizontal flip0.5, brightness/contrast0.15 and saturation0.1, without random crops. Float32 parameters with BF16 autocast train on GPU; final selection/calibration and held-out inference use float32 with TF32 disabled. Zero-shot uses a text direction on the endpoint embedding rather than the supervised temporal pooling, so that transition changes supervision and representation.",
        font=11.5,
        row_h=38,
    )
    table(
        "Training / changed parameters and checkpoints",
        "Six actual visual fine-tuning runs were completed",
        "Primary seed 42; seeds 43 and 44 repeat the selected visual learning rate",
        [
            "Encoder",
            "Visual params",
            "Head params",
            "Probe C",
            "Visual LR",
            "Best epoch",
        ],
        [
            [
                SHORT[m],
                f"{e.params[m]['vision_parameters']:,}",
                f"{e.params[m]['head_parameters']:,}",
                e.probes[m]["C"],
                e.selection[m, 42]["finetuned_lr"],
                e.selection[m, 42]["finetuned_epoch"],
            ]
            for m in MODELS
        ],
        [185, 155, 140, 110, 140, 146],
        "Primary MobileCLIP chooses epoch 3; SigLIP chooses epoch 2. Later epochs do not replace them on test.",
        "Both primary models select visual learning rate3e-5 from1e-5/3e-5. Epochs maximize development macro AP, averaging source AP within each of the five targets. Checkpoints and per-task routes are frozen before final test. Original AMP diagnostics are preserved; a numerical audit found batch-shape sensitivity, so the final tables use the separate float32_evaluation outputs. Operating thresholds were recalibrated only on float32 development predictions, keeping weights and method selection unchanged.",
        row_h=50,
    )
    for model in MODELS:
        stage_chart(e, model)
        for seed in (42, 43, 44):
            for method in METHODS:
                for task, source in [(t, None) for t in TASKS] + [
                    ("smoking", "stanford40")
                ]:
                    split = "external" if source else "test"
                    cell = e.clip[model, seed][method][split][task]
                    for point in POINTS:
                        export_metric(
                            "clip_visual_15oct",
                            model,
                            method,
                            task,
                            source or SOURCES[task],
                            point,
                            e.cm(model, method, task, point, seed, source),
                            seed=seed,
                            threshold=cell["thresholds"][point],
                        )
    for task, source, title in [(t, None, NAMES[t]) for t in TASKS] + [
        ("smoking", "stanford40", "Smoking / Stanford40 photos")
    ]:
        values = [
            metric_values(
                SHORT[model],
                METHOD_NAMES[method],
                e.cm(model, method, task, source=source),
            )
            for model in MODELS
            for method in METHODS
        ]
        m = e.cm(MODELS[0], "zero_shot", task, source=source)
        metric_table(
            title + ": zero-shot → probe → visual tune",
            f"Same held-out cohort: n={m['n']:,}; positives={m['positive']:,}; negatives={m['negative']:,} • dev-max-F1 cutoffs",
            values,
            "All stages are shown, including regressions. These are candidate results, not test-selected routes.",
            "Primary seed42. AP and ROC-AUC are ranking metrics independent of the displayed cutoff. F1, precision, recall and FPR use each method's own frozen maximum-development-F1 threshold. Zero-shot/probe/visual-tune stages use the same held-out labels; the zero-shot endpoint and supervised pooling representations differ. Dataset/source qualifications apply: historical CCTV and unattended labels are retrospective; IWDD is whole-clip waste dumping; loitering is one-camera supervised windows; sleep is18 provisional photos with6 positives; Stanford40 is an official photo holdout from a domain contributing training examples. Test-best method values do not override development routing.",
            warn=True,
        )
    table(
        "Evaluation / development-selected primary routes",
        "The chosen route is not always the most fine-tuned model",
        "Selection uses development mean source AP per target; simpler methods win ties",
        ["Target", "MobileCLIP route", "SigLIP2 route", "Mobile AP", "SigLIP AP"],
        [
            [
                NAMES[t],
                METHOD_NAMES[e.selection[MODELS[0], 42]["task_routes"][t]],
                METHOD_NAMES[e.selection[MODELS[1], 42]["task_routes"][t]],
                f(
                    e.cm(MODELS[0], e.selection[MODELS[0], 42]["task_routes"][t], t)[
                        "ap"
                    ]
                ),
                f(
                    e.cm(MODELS[1], e.selection[MODELS[1], 42]["task_routes"][t], t)[
                        "ap"
                    ]
                ),
            ]
            for t in TASKS
        ],
        [240, 180, 180, 138, 138],
        "SigLIP keeps the probe for dumping/sleep and zero-shot for loitering, even where a test alternative scores higher.",
        "Primary MobileCLIP routes all five targets through its fine-tuned checkpoint. Primary SigLIP routes smoking/unattended through visual tuning, dumping/public sleep through its probe and loitering through zero-shot. SigLIP's sleep fine-tune has a higher test AP than the selected probe; that fact is retained rather than retrospectively replacing the development route. These scores remain on the original 7,498-example study; later expanded-data strategies are a separate comparison.",
        row_h=42,
    )
    values = []
    for model in MODELS:
        method = e.selection[model, 42]["task_routes"]["unattended_items"]
        for point, label in (
            ("raw", "Raw zero"),
            ("dev_max_f1", "Dev max-F1"),
            ("dev_fpr_budget", "Dev ≤5% FPR"),
        ):
            values.append(
                metric_values(
                    SHORT[model], label, e.cm(model, method, "unattended_items", point)
                )
            )
    metric_table(
        "Unattended items: the cutoff changes the conclusion",
        "Primary development-selected routes • same 962 endpoints, 242 positive / 720 normal",
        values,
        "SigLIP raw F1 ≈0.749 and FPR ≈0.7%; its development-selected cutoffs produce ≈99% test FPR.",
        "No test threshold is fitted. The raw-zero cutoff was predeclared, so reporting it alongside development-max-F1 and development-budget results is a sensitivity analysis rather than threshold cherry-picking. AP is unchanged across all three rows for a route. Normal endpoint FPR remains different from false alerts per continuous camera-hour. Item ownership and temporal consistency are not measured by high endpoint AP.",
        warn=True,
    )
    seed_rows = []
    for task in TASKS:
        row = [NAMES[task]]
        for model in MODELS:
            scores = [
                e.cm(model, e.selection[model, s]["task_routes"][task], task, seed=s)[
                    "ap"
                ]
                for s in (42, 43, 44)
            ]
            row += [f(scores[0]), f"{min(scores):.3f}–{max(scores):.3f}"]
        seed_rows.append(row)
    table(
        "Evaluation / repeated optimization seeds",
        "Keep seed 42 primary and report the other seeds",
        "Each seed has its development-selected target route • same test sources in every repetition",
        ["Target", "Mobile primary", "Mobile range", "SigLIP primary", "SigLIP range"],
        seed_rows,
        [240, 159, 159, 159, 159],
        "These ranges measure optimization and route variation, not confidence intervals over unseen cameras.",
        "The range includes seeds42/43/44; each replication follows its own development-only routing and the primary selected visual LR. Seeds do not add independent cameras, people, new labels or events. The full CSV contains every method, seed and threshold. No winning test seed replaces seed42. Small public-sleep and loitering cohorts make apparent score changes especially unstable.",
        row_h=42,
    )


def strategy_and_matched_comparison(e):
    table(
        "Training / expanded representation study",
        "After backbone training: regions, geometry and new heads",
        "16 Oct expanded corpus • 330 configurations = 320 classifier fits + 10 text references",
        ["Candidate family", "What changes", "What stays controlled"],
        [
            [
                "Scene / predicted crop",
                "Full scene plus contextual person/item crop",
                "No source boxes select test crops",
            ],
            [
                "Ordered / geometry",
                "Chronological views; sparse automatic box features",
                "Timing and source splits retained",
            ],
            [
                "Weight interpolation",
                "Pretrained ↔ earlier tuned visual weights",
                "Alpha = 0, .25, .5, 1; selected on dev",
            ],
            [
                "Linear / MLP classifier",
                "LBFGS probes or 64-unit ReLU / dropout .3",
                "Same expanded training labels; source balancing",
            ],
            [
                "Selection",
                "Per-target mean development-source AP",
                "Primary seed42; no best-test route replacement",
            ],
        ],
        [200, 340, 336],
        "Fitting a new classifier or blending saved weights is distinct from a new visual-backbone fine-tune.",
        "The expanded training/development pool is3,973/1,812. Automatic Faster R-CNN contextual crops have confidence0.5,20% padding and full-frame fallback. Full pooled features use last/mean/max/change; ordered features concatenate four indexed frames plus mean/max. MLP uses one64-unit ReLU layer, dropout0.3, AdamW LR0.001/weight decay0.01, batch128, at most80epochs and patience8. Linear classifiers use LBFGS up to400iterations. Visual interpolation uses only previously trained parameters. Primary selection maximizes mean development-source AP. Secondary worst-source analyses and later policies have their own timing qualifications.",
        font=12,
        row_h=43,
    )
    cohorts = [(t, SOURCES[t], NAMES[t]) for t in TASKS]
    extras = [
        ("smoking", "kaggle_sujaykapadnis_smoking", "Smoking / Kaggle photos"),
        ("smoking", "stanford40", "Smoking / Stanford40"),
        ("smoking", "hmdb51_action_transfer", "Smoking / HMDB action"),
        ("rough_sleeping", "commons_new_author_transfer", "Sleep / new authors"),
        ("unattended_items", "meva_abandon_action_transfer", "MEVA abandon action"),
    ]
    total = 0
    for task, source, label in cohorts + extras:
        sm = e.sm(task, source)
        total += sm["n"]
        for model in MODELS:
            for point in POINTS:
                value = e.st(model, task, source, point)
                assert (value["n"], value["positive"], value["negative"]) == (
                    sm["n"],
                    sm["positive"],
                    sm["negative"],
                )
                threshold = e.strategy_routes[model]["tasks"][task]["thresholds"][point]
                export_metric(
                    "strategy_16oct",
                    model,
                    "dev_selected_strategy",
                    task,
                    source,
                    point,
                    value,
                    threshold=threshold,
                )
    assert total == e.provenance["exact_heldout_labels_checked"] == 3792
    for group, title in (
        (cohorts, "Five-target comparison on identical held-out cases"),
        (extras, "Photo and action-transfer comparison on matched cases"),
    ):
        values = []
        for task, source, label in group:
            sm = e.sm(task, source)
            values.append(
                [
                    label,
                    sm["n"],
                    sm["positive"],
                    f(e.st(MODELS[0], task, source)["ap"]),
                    f(e.st(MODELS[1], task, source)["ap"]),
                    f(sm["ap"]),
                ]
            )
        table(
            "Evaluation / matched three-model pipelines",
            title,
            "16 Oct CLIP strategies vs 19 Oct SmolVLM routed • same expanded train/dev and held-out labels",
            ["Cohort", "n", "Positive", "Mobile AP", "SigLIP AP", "Smol AP"],
            values,
            [286, 80, 95, 138, 138, 139],
            "This is a pipeline comparison: encoders, crops, resolution and adaptation methods differ.",
            "Existing provenance checked exact held-out IDs/labels for all3,792 cases. The new presentation also asserts all cohort denominators and class counts match across the three saved pipelines. Primary development-selected CLIP strategy routes are compared with primary seed42 SmolVLM per-target routes. The SmolVLM original control remains the causal comparison for LoRA weight adaptation. No VadCLIP score is inserted because its separate1,056-frame smoking test has different sessions. Later HMDB-adapted205-clip results use additional training and do not replace the906-clip transfer row. These historical cohorts remain retrospective.",
            row_h=42,
        )
    values = []
    for task in ("smoking", "unattended_items"):
        source = SOURCES[task]
        for model in MODELS:
            values.append(
                metric_values(
                    SHORT[model],
                    NAMES[task].replace("Smoking / ", ""),
                    e.st(model, task, source, "dev_fpr_budget"),
                )
            )
        values.append(
            metric_values(
                "SmolVLM",
                NAMES[task].replace("Smoking / ", ""),
                e.sm(task, source, point="dev_fpr_budget"),
            )
        )
    metric_table(
        "Matched cohorts: ranking gains versus operating errors",
        "16 Oct selected CLIP strategies / 19 Oct routed SmolVLM • development-budget cutoffs",
        values,
        "High unattended AP still coexists with high normal FPR. The development 5% budget does not transfer.",
        "Smoking cohort646/393positive; unattended962/242positive. Method-specific development-budget thresholds are fixed. These older studies constrain pooled development negatives for each target, rather than every source separately. The table includes AP, ROC-AUC, F1, precision, recall and FPR on exactly matched held-out labels. No low-FPR operating threshold is fitted to these test cases. Confusion counts and thresholds are retained in metrics.csv; precision and F1 can look favorable with high positive prevalence even when normal rejection fails.",
        warn=True,
    )


def smol_training_and_evaluation(e):
    table(
        "Training / SmolVLM language adaptation",
        "LoRA adapts the answer function while vision stays frozen",
        "SmolVLM-500M-Instruct • 3,973 training examples • five task-conditioned binary queries",
        ["Setting", "Executed configuration"],
        [
            [
                "Updated weights",
                "4,341,760 language-projection parameters; q/k/v/o and gate/up/down",
            ],
            [
                "LoRA setup",
                "Rank8; alpha16; dropout0.05; vision/connector/embeddings/output head frozen",
            ],
            [
                "Objective",
                "Full-vocabulary first-answer-token CE for Yes / No; mask every input token",
            ],
            [
                "Optimizer",
                "AdamW; LR3e-5 or1e-4; weight decay .01; warmup5%; gradient clip1.0",
            ],
            [
                "Effective batch",
                "Microbatch8 × accumulation4 =32 examples; max6epochs; patience2",
            ],
            [
                "Sampling / precision",
                "Each train case once + equal balanced resample; BF16 train, float32 eval",
            ],
        ],
        [205, 671],
        "No SmolVLM linear probe or full-parameter fine-tune was run; LoRA is the actual weight-adaptation stage.",
        "A sample contributes supervision only for its annotated target. The Yes/No first-token loss uses the whole vocabulary, with all prompt/image loss positions excluded. Vision tower, multimodal connector, token embeddings and vocabulary head are immutable. Rank8alpha16dropout0.05 target seven projection types in language layers. Linear LR decay follows5%warmup; gradient norm clipping1.0. One epoch has7,946 visits: all3,973unique examples plus an equally sized task/class/source/group-balanced resample. Frozen vision tokens are512-pixel inputs without image splitting; static tasks1image, temporal windows4, littering8. Held-out logits are float32 with TF32 off. Original and adapted routes use identical prompts and images.",
        font=12,
        row_h=38,
    )
    history_rows = []
    for route, history in e.smol_histories.items():
        best = max(history, key=lambda x: x["dev_macro_ap"])
        history_rows.append(
            [
                route.replace("seed_", "s").replace("/lr_", " / "),
                len(history),
                best["epoch"],
                f(best["dev_macro_ap"]),
                f(history[-1]["loss"]),
            ]
        )
    table(
        "Training / SmolVLM checkpoint selection",
        "All three adapter runs select epoch one",
        "Primary seed42 compares two learning rates; seed43 repeats only the primary winning LR",
        [
            "Run / visual tower frozen",
            "Epochs run",
            "Best epoch",
            "Best dev AP",
            "Last train loss",
        ],
        history_rows,
        [300, 130, 130, 158, 158],
        "More completed epochs do not imply a better selected model. Checkpoints maximize development macro source AP.",
        "Development metric averages dataset-source AP within each target, then equally across five targets. Original matched development macro AP0.5657; primary joint adapter seed42/LR1e-4 reaches0.7026. The low-rate candidate reaches0.6584; repeated seed43/LR1e-4 reaches0.6548. Each run completes three epochs under patience2 and selects epoch1. Epoch loss is not directly comparable to CLIP's binary classifier loss. Per-task routing chooses between original and the two primary adapters using development data before held-out inference; it is not a newly trained third head.",
        row_h=49,
    )
    cohorts = [(t, SOURCES[t], NAMES[t]) for t in TASKS] + [
        ("smoking", "stanford40", "Smoking / Stanford40")
    ]
    extra = [
        ("smoking", "kaggle_sujaykapadnis_smoking", "Smoking / Kaggle photos"),
        ("smoking", "hmdb51_action_transfer", "Smoking / HMDB action"),
        ("rough_sleeping", "commons_new_author_transfer", "Sleep / new authors"),
        ("unattended_items", "meva_abandon_action_transfer", "MEVA abandon action"),
    ]
    for task, source, label in cohorts + extra:
        for seed in (42, 43):
            sel = e.smol_selection[seed]
            for method in ("original", "joint", "routed"):
                route = (
                    "original"
                    if method == "original"
                    else sel["joint_adapter"]
                    if method == "joint"
                    else sel["task_routes"][task]
                )
                for point in POINTS:
                    export_metric(
                        "smol_lora_19oct",
                        "smolvlm-500m",
                        method,
                        task,
                        source,
                        point,
                        e.sm(task, source, method, point, seed),
                        seed=seed,
                        threshold=sel["thresholds"][route][task][point],
                    )
    values = []
    for task, source, label in cohorts:
        sm = e.sm(task, source)
        values.append(
            [
                label,
                sm["n"],
                f(e.sm(task, source, "original")["ap"]),
                f(e.sm(task, source, "joint")["ap"]),
                f(sm["ap"]),
                "LoRA low LR" if task == "rough_sleeping" else "LoRA high LR",
            ]
        )
    table(
        "Evaluation / SmolVLM stages",
        "Original model → one joint LoRA → per-target routing",
        "Same images, prompt and first-answer score within SmolVLM • primary seed42",
        ["Cohort", "n", "Original AP", "Joint AP", "Routed AP", "Routed weights"],
        values,
        [245, 70, 135, 125, 125, 176],
        "Routing improves the public-sleep choice; loitering ranking still falls below the original model.",
        "Original is the unadapted model under the same binary query, not the earlier six-label JSON generation baseline. Joint uses high-rate seed42adapter for all targets; primary routing selects low-rate seed42 for public sleep and high-rate for the other four. Loitering original AP0.4250 falls to0.1822 despite favorable primary development selection. Public sleep joint0.6327 versus routed0.7885 demonstrates why a joint model and a routed system must not be conflated. This table uses exact stored AP values rather than generated class labels.",
        font=12,
        row_h=37,
    )
    values = [
        metric_values(NAMES[t], "Routed LoRA", e.sm(t, SOURCES[t])) for t in TASKS
    ]
    values.append(
        metric_values("Stanford40", "Routed LoRA", e.sm("smoking", "stanford40"))
    )
    metric_table(
        "SmolVLM: full decision metrics after adaptation",
        "Primary per-target routes • maximum-development-F1 cutoffs • all five target cohorts",
        values,
        "Full-vocabulary Yes-minus-No margins support ranking; these are not calibrated event probabilities.",
        "All six rows use the frozen maximum-development-F1 operating point. The source-specific positive/negative counts are in the dataset and exported metrics tables. This binary task query has no trained unknown category, explanations or boxes. It is not numerically comparable to the older six-class free-generation accuracy. Native image forwards validate the cached visual-token path. Low-FPR and raw-zero operating points follow on the next slide; they do not use test fitting.",
        warn=True,
    )
    values = []
    for task in ("smoking", "unattended_items"):
        for point, label in (
            ("raw", "Raw zero"),
            ("dev_max_f1", "Dev max-F1"),
            ("dev_fpr_budget", "Dev ≤5% FPR"),
        ):
            values.append(
                metric_values(
                    NAMES[task], label, e.sm(task, SOURCES[task], point=point)
                )
            )
    metric_table(
        "SmolVLM: high endpoint recall can mean many false positives",
        "Primary routed model • compare three frozen operating points on the same cases",
        values,
        "Unattended at the dev-budget cutoff: 229/242 positives found, but 510/720 normals flagged.",
        "At the pooled-development-budget cutoff, unattended recall94.6%,FPR70.8%. Smoking CCTV recall1.8%,FPR0%. At raw zero, unattended recall64.9%,FPR25.8%; CCTV predicts no positives. These outcomes illustrate different score distributions and source offsets. Test cutoffs are unchanged. All three points share the same AP/AUC because only the threshold varies. Neither static-frame accuracy nor sparse-window AP measures persistent owner departure or long normal camera exposure.",
        warn=True,
    )
    values = []
    for task, source, label in [(t, SOURCES[t], NAMES[t]) for t in TASKS]:
        values.append(
            [
                label,
                f(e.sm(task, source, "original")["ap"]),
                f(e.sm(task, source, "joint", seed=42)["ap"]),
                f(e.sm(task, source, "joint", seed=43)["ap"]),
            ]
        )
    table(
        "Evaluation / SmolVLM seed replication",
        "A second seed exposes unstable task gains",
        "Same high-rate joint architecture and source cohort • primary route remains seed42",
        ["Cohort", "Original AP", "Joint seed42", "Joint seed43"],
        values,
        [291, 195, 195, 195],
        "Loitering regresses in both joint seeds; public-sleep AP varies strongly on only six positive test photos.",
        "The second seed is an optimization repeat, not an independent data-source replication. Seed42joint sleep AP0.6327,seed43joint0.9167 and primary routed low-rate0.7885 are different configurations. Optional repeated-gain routing was formulated after primary held-out inspection; it is exploratory and not substituted into these primary tables. Exact original baseline rows repeat across seeds only as controls, not as new model runs.",
        row_h=42,
    )


def vad_training_and_evaluation(e):
    paired(
        "Training / VadCLIP temporal adapter",
        "VadCLIP learns a sequence model on frozen CLIP-B/16",
        "Corrected offline supervised smoking study • 1,910 train / 312 dev / 1,056 test frames",
        (
            "Frozen-feature controls",
            "Two fixed text prompts\nCLIP-B/16 image embeddings\nLogistic probe C1, max2,000steps\nSame training/test frames as adapter\nDevelopment-only F1 cutoffs",
        ),
        (
            "Actual adapter training",
            "About12.642M trainable parameters\nTwo temporal layers + graph branches\nLength256; attention window8\nAdamW LR.001, decay.01; batch4\nSeed234; max100epochs, patience50",
        ),
        "This is offline sequence-context evaluation. It does not establish a causal five-anomaly event detector.",
        "The pretrained CLIP-B/16 feature/text encoder stays frozen. The baseline logistic regression uses C1,max_iter2000 and training-frame labels, without test tuning. The adapter uses masked per-frame binary visual loss plus alignment cross entropy. Padded frames do not contribute. It can use context inside an offline sequence; raw-frame scores are expanded from the representative sequence to original labels. Best checkpoint/branch uses development PR performance, selecting epoch1; training stops after51completed epochs under patience50. Test threshold comes only from312development frames. This session split shares physical cameras across dates and differs from the other models' unseen-camera holdout. Historical earlier adapter ablations with no final test are excluded.",
    )
    values = []
    for stage, label, metric in (
        ("zero_shot", "CLIP text", e.vad_base["metrics"]["zeroshot"]),
        ("linear_probe", "Linear probe", e.vad_base["metrics"]["linearprobe"]),
        ("temporal_adapter", "Temporal adapter", e.vad_eval["metrics"]["dual_mean"]),
    ):
        m = export_metric(
            "vadclip_session_13oct",
            "vadclip",
            stage,
            "smoking",
            "cctv_session_test",
            "dev_max_f1",
            metric,
            seed=234,
            threshold=metric["threshold"],
        )
        values.append(metric_values("VadCLIP", label, m))
    constant = {
        "ap": 560 / 1056,
        "auroc": 0.5,
        "f1": 1120 / 1616,
        "precision": 560 / 1056,
        "recall": 1,
        "fpr": 1,
    }
    values.append(metric_values("Reference", "Always positive", constant))
    metric_table(
        "VadCLIP: zero-shot → probe → temporal adapter",
        "Same 1,056 held-out frames: 560 positive / 496 normal • per-stage dev-max-F1 cutoffs",
        values,
        "Adapter AP rises from0.482 to0.633, but FPR remains35.7% and recall51.8% at its frozen cutoff.",
        "The text baseline AP0.4823,linear probe0.5439,dual-mean adapter0.6332 all use identical test labels. The adapter's ROC-AUC0.5874,F10.5648,precision0.6210,recall0.5179,FPR0.3569,accuracy0.5767. Exact confusion:290TP,177FP,270FN,319TN. The probe's development cutoff flags97.8%of test normals, despite its high recall. Always-positive F1≈0.693 exceeds adapter F1 while producing100%FPR, demonstrating why F1 alone is misleading. AP/AUPR are taken from the saved implementation; this is not a matched-cohort comparison with646-case camera-c1 tests.",
        warn=True,
    )
    values = []
    for model in MODELS:
        values.append(
            metric_values(
                SHORT[model] + " • A",
                "Visual tune",
                e.cm(model, "finetuned", "smoking"),
            )
        )
    values.append(
        metric_values(
            "SmolVLM • A", "Joint LoRA", e.sm("smoking", "smoking_cctv", "joint")
        )
    )
    v = e.vad_eval["metrics"]["dual_mean"]
    m = {
        "ap": v["aupr"],
        "auroc": v["auroc"],
        "f1": v["f1"],
        "precision": v["precision"],
        "recall": v["recall"],
        "fpr": v["false_positive_rate"],
    }
    values.append(metric_values("VadCLIP • B", "Temporal adapter", m))
    metric_table(
        "Four adapted models: metrics with their protocol labels",
        "A: 646-frame camera holdout (393+) • B: 1,056-frame session holdout (560+) • dev-max-F1",
        values,
        "Metric columns are shared; supervision, training size and test split are not equal across all four rows.",
        "A includes three pipelines sharing held-out camera-c1 frames but different primary training corpora and score interfaces. MobileCLIP/SigLIP entries are the actual visual-tune candidate rather than a later strategy selected after source-data expansion. SmolVLM is the high-rate primary joint LoRA. B is the corrected separate VadCLIP temporal adapter test. Do not infer a winner by sorting AP or F1 across A/B; prevalence is60.8%versus53.0%,and split difficulty differs. The preceding within-study stage tables are the valid attribution comparisons. The expanded-data matched three-model table uses a different, explicitly named strategy stage.",
        warn=True,
    )


def later_experiments_and_summary(e):
    table(
        "Training / person and temporal heads",
        "Later learning adds automatic regions and temporal memory",
        "16–17 Oct CLIP pilot • frozen pretrained visual encoders • 50 saved candidate heads",
        ["Component", "Executed training / representation", "Evaluation boundary"],
        [
            [
                "Input evidence",
                "16 HMDB views; full / body / upper-body crops",
                "CCTV smoking test remains static",
            ],
            [
                "Temporal classifier",
                "128-unit projection;64-unit causal GRU; optional18geometry channels",
                "Past-only sequence features; no oracle tracks",
            ],
            [
                "Person supervision",
                "1,460 matched train regions,277positive; scene BCE + .5region BCE",
                "Automatic test regions; no test box selection",
            ],
            [
                "Optimization",
                "AdamW LR.0005,decay.01; batch128;max60epochs;patience8",
                "Worst dev-source AP, then mean AP",
            ],
            [
                "Additional source data",
                "530 train /161dev /205test HMDB clips after title-part purge",
                "Different cohort from906-clip transfer test",
            ],
        ],
        [180, 405, 291],
        "These are new classifiers and temporal heads; the visual encoders are frozen in this pilot.",
        "Each encoder fits14smoking and11unattended heads. Person/region baselines, causal GRUs with and without geometry, three seeds and region-supervised variants are retained. Hard negatives are mined only from training normals beginning epoch4; a checkpoint selected earlier did not benefit from later mining. Smoking train3,342/dev1,475 combines expanded static data and HMDB adaptation; unattended uses236/236eligible windows. Ten known film-part aliases were removed from training. Improved crop geometry does not prove improved anomaly classification, and sparse ownership geometry remains a hypothesis. This training stage differs from both the earlier partial visual fine-tunes and SmolVLM language LoRA.",
        font=11.5,
        row_h=43,
    )
    pilot = {m: load(PILOT / m / "results.json") for m in MODELS}
    pilot_selections = {m: load(PILOT / m / "joint_selection.json") for m in MODELS}
    latest17 = load(ROOT / "results/event_improvements_20261017/smoking_results.json")
    sel17 = load(ROOT / "results/event_improvements_20261017/smoking_selection.json")
    selected16 = {
        MODELS[0]: "person_supervised_hardneg_seed_42",
        MODELS[1]: "region_temporal_seed_42",
    }
    for source, label in (
        ("smoking_cctv", "CCTV frames, n646"),
        ("hmdb51_action_transfer", "Adapted HMDB test, n205"),
    ):
        values = []
        for model in MODELS:
            for name, stage in (
                ("scene_linear", "Scene probe"),
                (selected16[model], "Region / person"),
            ):
                m = pilot[model]["clip_metrics"]["smoking"][name][source][
                    "operating_points"
                ]["dev_fpr_budget"]
                values.append(metric_values(SHORT[model], stage, m))
                export_metric(
                    "video_pilot_16oct",
                    model,
                    name,
                    "smoking",
                    source,
                    "dev_fpr_budget",
                    m,
                    threshold=next(
                        c
                        for c in pilot_selections[model]["tasks"]["smoking"][
                            "candidates"
                        ]
                        if c["name"] == name
                    )["thresholds"]["dev_fpr_budget"],
                )
            choice = sel17[model]["selected"]
            m = latest17[model][choice]["sources"][source]
            values.append(metric_values(SHORT[model], "Sourcewise policy", m))
            export_metric(
                "source_policy_17oct",
                model,
                choice,
                "smoking",
                source,
                "dev_every_source_budget",
                m,
                threshold=latest17[model][choice]["threshold"],
            )
        metric_table(
            "Later smoking stages: " + label,
            "Within-pilot training/control cohort • original pooled budget vs later every-source policy",
            values,
            "Later policies change selection/calibration; the three rows do not isolate only a larger neural network.",
            "Scene probe and region/person rows use their original pooled development-budget cutoff. The17October policy chooses a fixed ensemble and single global cutoff that satisfies each development source's5%FPR budget. Primary policy uses person-supervised hard-negative seed42 for MobileCLIP and all-head ensemble for SigLIP. These policies are selected on development data, while historical test sources were already inspected. The205HMDB test has30positive/175controls,additional530training clips and16views; the earlier906clip Smol/strategy transfer cohort had no HMDB training and8views. Do not compare those APs as one fixed-test learning curve.",
            warn=True,
        )
    table(
        "Training / source-data ablation",
        "20 October: old-only, new-only and mixed training",
        "740 new scripted-action clips • no new visual-backbone or SmolVLM weight training",
        ["Setting", "Frozen experimental design"],
        [
            [
                "New labels",
                "Smoking gesture vs drinking/eating/phone;masked actors;no cigarette confirmation",
            ],
            [
                "Source split",
                "Train340/dev160/test240;scene blocks and numeric subject aliases disjoint",
            ],
            [
                "Classifier families",
                "Old-only2,812/new-only340/mixed3,152 training examples",
            ],
            [
                "Candidate budget",
                "2encoders ×3families ×2full/crop recipes ×3C values =36logistic fits",
            ],
            [
                "Selection and views",
                "Same1,474dev cases;worst-source recall at every-source≤5%FPR;16newviews",
            ],
            ["Held-out scope", "1,916cases per encoder:historical1,676 + new proxy240"],
        ],
        [195, 681],
        "More source data is tested as weak supervision. Scene4 has only three held-out subject groups.",
        "XRF55 official public RGB release v6 supplies740selected clips and11,840frames. Training-only inspection found masked actors performing hand-to-mouth gestures; class15 is retained as a publisher action proxy, not verified target smoking. New-only fitting uses340weakly labelled clips, not a verified CCTV smoking corpus. Frozen pretrained encoders and exact automatic contextual crops feed logistic heads with C0.1/1/10. Each source/class/group is balanced. All36fits converge; both model development selections freeze before test. No new temporal rule, visual fine-tune or language adapter is trained. Old-only also receives new-source development calibration, and new-only receives old-source development feedback.",
        font=11.5,
        row_h=38,
    )
    for source, label in (
        ("smoking_cctv", "Historical CCTV"),
        ("xrf55_rgb_public", "New action proxy"),
    ):
        values = []
        for model in MODELS:
            for family in ("old_only", "new_only", "mixed"):
                entry = e.latest[model]["results"][family + "/" + source]
                for point, m in entry["operating_points"].items():
                    export_metric(
                        "source_ablation_20oct",
                        model,
                        family,
                        "smoking",
                        source,
                        point,
                        m,
                        threshold=e.latest_selections[model]["candidates"][
                            entry["candidate"]
                        ]["thresholds"][point],
                    )
                values.append(
                    metric_values(
                        SHORT[model],
                        family.replace("_", " "),
                        entry["operating_points"]["dev_fpr_budget"],
                    )
                )
        metric_table(
            "Data-addition outcome: " + label,
            "Frozen per-family development selection • every-source budget cutoff • same source cohort in all rows",
            values,
            "SigLIP improves proxy AP but loses CCTV ranking and normal rejection; no new model is promoted.",
            "Mixed SigLIP CCTV AP0.7309to0.6227,and FPR3.95%to36.76%. On the new action proxy AP0.3568to0.7768,with recall21.7%and FPR1.7%at the frozen cutoff. MobileCLIP also loses CCTV AP0.6937to0.5673; its mixed proxy cutoff finds no positives and flags20drinking clips. All source control errors in these CLIP families occur on drinking actions. The new-source subject-level recall varies strongly; mixed SigLIP recalls10%,0%,55%across the three aliases. HistoricalCCTV remains retrospective, and proxy gain does not establish true CCTV improvement.",
            warn=True,
        )
    values = []
    for name in ("original", "lora"):
        result = e.latest_smol[name]
        for key, label in (
            ("new_source_calibrated", "New-source dev"),
            ("original19oct_cutoffs", "19 Oct dev"),
        ):
            m = result[key]["dev_fpr_budget"]
            values.append(metric_values("Smol " + name, label, m))
            export_metric(
                "source_transfer_20oct",
                "smolvlm-500m",
                name,
                "smoking",
                "xrf55_rgb_public",
                key + "_budget",
                m,
                threshold=e.latest_smol_calibration[name][
                    "new_source_thresholds"
                    if key == "new_source_calibrated"
                    else "original19oct_thresholds"
                ]["dev_fpr_budget"],
            )
    metric_table(
        "New-source SmolVLM transfer after prior LoRA",
        "No new training • original vs fixed prior smoking adapter •8views each;240test clips,60positive",
        values,
        "Original AP0.386 → LoRA0.357. Both older budget cutoffs miss all new-source positives.",
        "Both routes use the same8uniform frames selected from the16-view source manifest; the existing native prompt supports at most8. Each route scores160dev clips and freezes new-source thresholds before either scores test. Original new-dev-budget recall61.7%,FPR36.7%;LoRA35.0%,18.3%. Older19Octoberbudget cutoffs yield zero recall and zero FPR on both routes. No scene4 adapter selection or new LoRA training occurs. Native image/cache margin checks pass. The source remains scripted-gesture evidence and cannot certify real smoking performance.",
        warn=True,
    )
    table(
        "Evaluation / model coverage and conclusion",
        "What the four-model evidence supports for each target",
        "Compare complete source/metric records; missing supervision cannot be filled with a score",
        ["Target", "MobileCLIP / SigLIP2", "SmolVLM", "VadCLIP"],
        [
            [
                "Smoking",
                "Photo gains; CCTV/source regressions persist",
                "LoRA mixed CCTV gain;proxy transfer regresses",
                "Temporal adapter AP gain;separate sessions",
            ],
            [
                "Littering",
                "Visual/probe ranking improves on IWDD",
                "LoRA ranking improves on IWDD",
                "Not trained/evaluated",
            ],
            [
                "Loitering",
                "Later geometry helps historical ranking",
                "Both joint seeds regress",
                "Not trained/evaluated",
            ],
            [
                "Public sleep",
                "Strong small-photo diagnostic",
                "Highly seed/route dependent",
                "Not trained/evaluated",
            ],
            [
                "Unattended",
                "High endpoint AP;threshold transfer poor",
                "Recall gains with excessive FPR",
                "Not trained/evaluated",
            ],
        ],
        [145, 281, 260, 190],
        "No four-model winner across five real events is established by the current training studies.",
        "The most defensible comparisons are within-study baseline/probe/adaptation stages and expanded-data matched three-pipeline cohorts. Generalization is still limited by camera diversity, independent annotation and confusable negatives. Apparent public sleep remains a visual proxy, and static endpoints do not verify owner departure or excessive dwell duration. The separate VadCLIP study supports only smoking-session results. Full-parameter model unfreezing was not performed. No test-best seed, retrospective cutoff or optional repeated-gain routing is silently promoted. The new deck reports existing training/evaluation facts without another inference or latency experiment.",
        font=11.5,
        row_h=43,
    )
    paired(
        "Training / artifact guide",
        "Detailed numbers and training choices are available to review",
        "Editable slides + vector PDF + complete stage metrics + source identities",
        (
            "Slide companion files",
            "Unrounded metrics and confusion counts\nPer-stage cutoff / seed / source fields\nDataset counts with positive labels\nDetailed speaker notes per slide\nSource hashes and measured layout",
        ),
        (
            "What remains necessary",
            "Compatible independent event labels\nNew camera/day source blocks\nDrinking / eating / phone controls\nTrack duration and owner/item history\nFrozen source-aware acceptance tests",
        ),
        "Read stage comparisons alongside source coverage and FPR; more fine-tuning alone has not solved transfer.",
        "Companion directory: reports/slides/training_evaluation_20261020. The CSV exports stored AP,ROC-AUC,precision,recall,F1,FPR,TP/FP/FN/TN,as well as derived accuracy/balanced accuracy. The manifest binds every reporting input and builder. A separate validation checks source hashes, metrics against the original result cells, PowerPoint/PDF content and rendered page bounds. The8-page20Octoberreport and102-page19Octobertechnical history remain separate documents. This deck focuses on the four requested model families and their training/evaluation stages.",
    )


def content(e):
    # Filled by the following section functions; each section stays scoped to
    # the experiment whose training split and score interface it describes.
    for split in ("train", "dev", "test"):
        group = [r for r in e.new_rows if r["split"] == split]
        positive = sum(r["target"] for r in group)
        DATASETS.append(
            {
                "study": "xrf55_proxy_20oct",
                "task": "smoking",
                "split": split,
                "n": len(group),
                "positive": positive,
                "negative": len(group) - positive,
            }
        )
    for task in ("smoking", "rough_sleeping", "unattended_items"):
        group = [r for r in e.extra_rows if r["task"] == task]
        positive = sum(r["target"] for r in group)
        DATASETS.append(
            {
                "study": "additional_transfer_19oct",
                "task": task,
                "split": "external",
                "n": len(group),
                "positive": positive,
                "negative": len(group) - positive,
            }
        )
    scope_and_data(e)
    clip_training_and_stages(e)
    strategy_and_matched_comparison(e)
    smol_training_and_evaluation(e)
    vad_training_and_evaluation(e)
    later_experiments_and_summary(e)


def summary_slide(e):
    values = []
    for model in MODELS:
        m = e.cm(model, "finetuned", "smoking")
        values.append(
            [
                SHORT[model] + " • visual",
                3258,
                "A: 646",
                *(f(e.cm(model, method, "smoking")["ap"]) for method in METHODS),
                f(m["f1"]),
                p(m["fpr"]),
            ]
        )
    sm = e.sm("smoking", "smoking_cctv", "joint")
    values.append(
        [
            "SmolVLM • LoRA",
            3973,
            "A: 646",
            f(e.sm("smoking", "smoking_cctv", "original")["ap"]),
            "Not run",
            f(sm["ap"]),
            f(sm["f1"]),
            p(sm["fpr"]),
        ]
    )
    v = e.vad_eval["metrics"]["dual_mean"]
    values.append(
        [
            "VadCLIP • temporal",
            1910,
            "B: 1,056",
            f(e.vad_base["metrics"]["zeroshot"]["aupr"]),
            f(e.vad_base["metrics"]["linearprobe"]["aupr"]),
            f(v["aupr"]),
            f(v["f1"]),
            p(v["false_positive_rate"]),
        ]
    )
    table(
        "Training / one-slide overview",
        "Four-model training stages: smoking evaluation",
        "Stages use their own development selection • adapted F1/FPR use dev-max-F1 cutoffs",
        [
            "Model / adapted part",
            "Train n",
            "Test n",
            "Base AP",
            "Probe AP",
            "Adapt AP",
            "Adapt F1",
            "Adapt FPR",
        ],
        values,
        [190, 96, 110, 96, 96, 96, 96, 96],
        "A: camera-c1 test. B: separate session test. Training sizes/units differ; this is not a matched leaderboard.",
        "This is a compact index to the detailed46-slide training/evaluation deck, which covers all five objectives, complete source/count tables and AP,ROC-AUC,F1,precision,recall,FPR,accuracy/balanced-accuracy exports. Train n is all-target examples for CLIP/Smol and smoking frames for VadCLIP. Primary CLIP uses3,258train/1,633dev and partial visual tuning; Smol uses3,973train/1,812dev and languageLoRA. Their A smoking test shares646frames/393positive but training corpora differ. VadCLIP B uses1,056frames/560positive across3session-held-out test groups. No SmolVLM linear probe was run. The main deck provides identical-case three-model comparisons and within-study learning stages.",
        warn=True,
        font=11.5,
        row_h=47,
    )


def main():
    d.SLIDES.clear()
    d.SOURCES.clear()
    EXPORTS.clear()
    DATASETS.clear()
    DETAIL.mkdir(parents=True, exist_ok=True)
    evidence = Evidence()
    content(evidence)
    full_slides = list(d.SLIDES)
    d.SLIDES.clear()
    summary_slide(evidence)
    overview = list(d.SLIDES)
    d.SLIDES[:] = full_slides
    for i, slide in enumerate(d.SLIDES, 1):
        for item in slide["items"]:
            if item["kind"] != "text":
                continue
            if item["x"] == 882 and item["y"] == 512:
                item["value"] = f"{i:02d}"
            assert (
                len(item["value"].splitlines()) * item["leading"] <= item["h"] + 0.01
            ), (slide["title"], item)
    d.check_geometry(d.SLIDES)
    rl_config.canvas_basefontname = "DejaVu"
    d.render_pptx(d.SLIDES, OUT / (STEM + ".pptx"))
    presentation = Presentation(OUT / (STEM + ".pptx"))
    presentation.core_properties.title = (
        "Training and evaluation: MobileCLIP, SigLIP2, SmolVLM and VadCLIP"
    )
    presentation.core_properties.subject = "Datasets, zero-shot baselines, probes, visual adaptation, LoRA and temporal adapters"
    presentation.save(OUT / (STEM + ".pptx"))
    d.render_pdf(d.SLIDES, OUT / (STEM + ".pdf"))
    d.render_pptx(overview, OUT / "training_stages_summary_20261020.pptx")
    summary = Presentation(OUT / "training_stages_summary_20261020.pptx")
    summary.core_properties.title = "Four-model training-stage summary: smoking"
    summary.core_properties.subject = "Qualified overview; see the detailed training/evaluation deck for five-target metrics"
    summary.save(OUT / "training_stages_summary_20261020.pptx")
    d.render_pdf(overview, OUT / "training_stages_summary_20261020.pdf")
    for path, records in (
        (DETAIL / "metrics.csv", EXPORTS),
        (DETAIL / "dataset_counts.csv", DATASETS),
    ):
        with path.open("w") as file:
            writer = csv.DictWriter(file, fieldnames=list(records[0]))
            writer.writeheader()
            writer.writerows(records)
    (DETAIL / "speaker_notes.md").write_text(
        "\n\n".join(
            ["# Four-model training and evaluation"]
            + [
                f"## {i:02d}. {s['title']}\n\n{s['notes']}"
                for i, s in enumerate(d.SLIDES, 1)
            ]
        )
        + "\n"
    )
    helpers = [Path(__file__), OUT / "build_presentation.py"]
    (DETAIL / "manifest.json").write_text(
        json.dumps(
            {
                "date": "2026-10-20",
                "slide_count": len(d.SLIDES),
                "models": [
                    "mobileclip2-s2",
                    "siglip2-b16-224",
                    "smolvlm-500m",
                    "vadclip",
                ],
                "sources_sha256": d.SOURCES,
                "builders_sha256": {
                    str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
                    for p in helpers
                },
                "metric_rows": len(EXPORTS),
                "dataset_rows": len(DATASETS),
                "layout": d.SLIDES,
                "summary_layout": overview,
                "no_new_model_training_or_inference": True,
                "no_latency_testing": True,
                "qualification": "Within-study stages and matched three-model pipelines are distinguished from separate VadCLIP smoking protocol; unavailable stages are explicit.",
            },
            indent=2,
        )
        + "\n"
    )
    print(
        "Created",
        len(d.SLIDES),
        "training/evaluation slides;",
        len(EXPORTS),
        "metric rows",
    )


if __name__ == "__main__":
    main()
