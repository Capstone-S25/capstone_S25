"""Build the complete accuracy/training progress deck from preserved artifacts."""

from __future__ import annotations

import hashlib
import json
from copy import deepcopy
from pathlib import Path

from pptx import Presentation

from reports.slides import build_presentation as d

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
RUN = ROOT / "results/event_improvements_20261017"
STEM = "model_evaluation_progress_20261017"
MODELS = ["mobileclip2-s2", "siglip2-b16-224"]
SHORT = {MODELS[0]: "MobileCLIP", MODELS[1]: "SigLIP"}


def register(path):
    path = ROOT / path
    d.SOURCES[str(path.relative_to(ROOT))] = hashlib.sha256(
        path.read_bytes()
    ).hexdigest()


def load(path):
    return d.read(ROOT / path)


def new(category, title, subtitle, source, notes):
    s = d.new(category, title, subtitle, source, notes)
    for item in s["items"]:
        if item["kind"] == "text" and "EVALUATION 13 OCT 2026" in item["value"]:
            item["value"] = item["value"].replace(
                "EVALUATION 13 OCT 2026", "PROGRESS 17 OCT 2026"
            )
    return s


def brief(s, value, color=d.TEAL, y=431):
    d.banner(s, y, value, color=color, size=14, h=61)


def paired(category, title, subtitle, source, left, right, takeaway, notes, warn=False):
    s = new(category, title, subtitle, source, notes)
    for x, (heading, body), color in [(42, left, d.TEAL), (491, right, d.ORANGE)]:
        d.card(s, x, 166, 427, 249, heading, body, color=color, size=16)
    brief(s, takeaway, d.ORANGE if warn else d.TEAL)
    return s


def tab(
    category,
    title,
    subtitle,
    source,
    headings,
    rows,
    widths,
    takeaway,
    notes,
    row_h=43,
    font=14,
    warn=False,
):
    s = new(category, title, subtitle, source, notes)
    d.table(s, 42, 170, widths, headings, rows, row_h=row_h, font=font, header_h=40)
    brief(s, takeaway, d.ORANGE if warn else d.TEAL)
    return s


def build_content():
    d.SLIDES.clear()
    d.SOURCES.clear()
    quality = load("results/expanded_accuracy_20261013/quality_summary.json")
    fine = load("results/clip_finetuning_20261015/verification.json")
    strategy = load("results/strategy_study_20261016/verification.json")
    pilot = load("results/video_pilot_20261016/verification.json")
    latest = load("results/event_improvements_20261017/verification.json")
    validation = load("results/event_improvements_20261017/validation.json")
    selection = load("results/event_improvements_20261017/smoking_selection.json")
    smoke = load("results/event_improvements_20261017/smoking_results.json")
    events = load("results/event_improvements_20261017/event_results.json")
    component = load("results/event_improvements_20261017/component_results.json")
    component_protocol = load(
        "data/event_improvements_20261017/component_protocol.json"
    )
    vad = load("results/gpu_all_models_20261013/vadclip/adapter/eval.json")
    assert (
        quality["status"]
        == strategy["status"]
        == pilot["status"]
        == latest["status"]
        == validation["status"]
        == "complete"
    )
    assert fine["study_status"] == "complete"
    for p in [
        "reports/workflow_review_20261013.md",
        "reports/expanded_accuracy_review_20261014.md",
        "reports/clip_finetuning_findings_20261016.md",
        "reports/strategy_findings_20261016.md",
        "reports/video_pilot_findings_20261016.md",
        "reports/event_improvements_findings_20261017.md",
        "results/clip_finetuning_20261015/comparison.csv",
        "results/strategy_study_20261016/comparison.csv",
        "results/video_pilot_20261016/comparison.csv",
        "results/video_pilot_20261016/events.csv",
        "results/event_improvements_20261017/protocol.json",
        "data/event_improvements_20261017/receipt.json",
    ]:
        register(p)

    s = new(
        "Findings / executive summary",
        "Five-anomaly model evaluation",
        "Complete evaluation and improvement history • 13–17 October 2026",
        "Preserved studies + 17 Oct metric replay",
        "This deck replaces neither historical results nor their selection protocols. It consolidates the four-model comparison, expanded data, six actual visual fine-tunes, 330 strategy configurations, 50 video-pilot heads and 38 new ensemble/threshold policies, plus tracker and tiled-detector ablations. Candidate/prediction counts reuse sources and are not independent events. No deployment winner across all five objectives is established. No new latency benchmark was run. Stage-specific regressions and data limits remain visible.",
    )
    d.box(s, 42, 166, 876, 101, d.NAVY)
    d.text(
        s,
        61,
        181,
        834,
        76,
        "Recognition improved in selected domains.\nReliable camera transfer is still unresolved.",
        26,
        d.WHITE,
        True,
    )
    for x, value, label, detail in [
        (42, "6", "Visual fine-tunes", "2 CLIP encoders × 3 seeds"),
        (
            340,
            "330 + 50",
            "Strategy / video routes",
            "New heads and routes; different studies",
        ),
        (
            638,
            "38 + tiles",
            "Latest improvements",
            "Dev-only policies + causal state replay",
        ),
    ]:
        d.stat(s, x, 282, 280, 142, value, label, detail)
    brief(
        s,
        "Strongest next investment: verified event histories and matched normal footage on independent target cameras.",
        d.ORANGE,
    )

    tab(
        "Objective / definitions",
        "Five objectives require different evidence",
        "Existing dataset annotations and definitions remain the source of truth",
        "Training protocol + source annotation audits",
        ["Objective", "Evidence required", "Current qualification"],
        [
            [
                "Smoking",
                "Person + hand/object/action evidence",
                "Low-resolution CCTV differs from photos",
            ],
            [
                "Littering",
                "Disposal action and scene context",
                "IWDD includes larger waste dumping",
            ],
            [
                "Loitering",
                "Identified person and timed history",
                "A visible person alone is insufficient",
            ],
            [
                "Rough sleeping",
                "Persistent apparent public sleep",
                "Posture photos are provisional proxies",
            ],
            [
                "Unattended items",
                "Item persistence + owner separation",
                "30 s criterion belongs to AVSS source",
            ],
        ],
        [169, 356, 351],
        "Unknown history or missing labels stay unknown; other objectives are never automatically labelled normal.",
        "The user accepted existing definitions and annotations. There is no newly invented universal loitering threshold. Public-sleep labels describe visible sleep/bedding/context, never housing status. AVSS specifies stationary for 30 seconds and an absent owner; historical provisional windows and MEVA abandonment actions do not uniformly implement this criterion. Grounding boxes and source annotations are separate from inferred ownership.",
        row_h=43,
        font=13,
    )

    tab(
        "Models / comparison scope",
        "Four image models and a separate VadCLIP study",
        "All five were investigated; later adaptation focuses on the two CLIP encoders",
        "Expanded accuracy report; supervised VadCLIP eval.json",
        ["Model", "Evaluation / adaptation", "Scope of conclusions"],
        [
            [
                "MobileCLIP2-S2",
                "Zero-shot → probe → visual fine-tune",
                "All five; region / temporal pilots",
            ],
            [
                "SigLIP2-B/16-224",
                "Zero-shot → probe → visual fine-tune",
                "All five; region / temporal pilots",
            ],
            [
                "SmolVLM-500M-Instruct",
                "Deterministic prompted responses",
                "Expanded four-model image study",
            ],
            [
                "MiniCPM-V 4.0",
                "Unquantized prompted responses",
                "Expanded four-model image study",
            ],
            [
                "VadCLIP",
                "Separate supervised smoking adapter",
                "Offline protocol; different supervision",
            ],
        ],
        [204, 342, 330],
        "VLMs were not fine-tuned in later rounds. Head fitting and visual-backbone updates are reported separately.",
        "Original model names: MobileCLIP2-S2, SigLIP2-B/16-224, SmolVLM-500M-Instruct, MiniCPM-V 4.0 and VadCLIP. Initial comparisons use identical inputs within each image protocol. VadCLIP receives different offline sequence supervision and cannot be ranked directly against the four image models as an equal-supervision five-task benchmark. Later encoder studies do not rerun the VLMs and must not imply their current architecture cannot improve.",
        row_h=43,
        font=13,
    )

    tab(
        "Work completed / chronology",
        "The work progressed from prompts to event evidence",
        "These counts describe separate experiments, often reusing the same source footage",
        "Completed run / verification artifacts for each stage",
        ["Stage", "Completed work", "Evidence produced"],
        [
            [
                "13 Oct",
                "4 models; 13,400 predictions",
                "Initial audit and VadCLIP branch",
            ],
            [
                "14 Oct",
                "55,156 expanded predictions",
                "Larger cohorts; qualified weak diagnostics",
            ],
            [
                "15–16 Oct",
                "6 actual CLIP visual fine-tunes",
                "Changed weights; three seeds each",
            ],
            [
                "16 Oct",
                "330 configurations; 50 pilot heads",
                "Data transfer, crops, temporal / tracks",
            ],
            [
                "17 Oct",
                "38 policies; state replay; 830 views",
                "Ensembles, memory, tiled item detections",
            ],
        ],
        [115, 363, 398],
        "New runs and artifacts preserve the earlier reports and slide decks.",
        "13 Oct: 3,350 protocol inputs per model, including shared endpoint variants; 104 input aliases removed. 14 Oct: 13,789 input/protocol evaluations per model and 40 summaries. Fine-tuning: 7,498 target-specific examples and 12,478 unique view images. Strategy screen: 330 configurations across five tasks and two backbones, including pretrained references; not 330 visual fine-tunes. Video pilot: 50 saved heads, 99,528 saved predictions. Latest: 19 ensemble/threshold policies per backbone, four item-rule ablations and 166 source-annotated images × five fixed views. Counts should not be added as independent events.",
        row_h=43,
        font=13,
    )

    paired(
        "Data / readiness",
        "Many files did not provide complete event labels",
        "The dataset audit changed what could be evaluated honestly",
        "Workflow review + expanded cohort definitions",
        (
            "Usable supervised evidence",
            "Smoking: CCTV person/action labels\nIWDD: dumping action clips + controls\nStreet Scene: person track/type labels\nCAVIAR: item/person source boxes\nABODA: reviewed time intervals",
        ),
        (
            "Useful diagnostics, weaker event truth",
            "6,667 littering candidate images\n1,675 person/loitering candidates\nSynthetic lying / falling postures\nStatic public-sleep photographs\nMEVA abandonment-action windows",
        ),
        "Positive-only images, augmented copies and pose labels cannot establish five-event precision or false-alarm rates.",
        "Weak collections were evaluated as target-prediction rates, not mislabeled as event accuracy. Littering exports name dropping/dumping/placing/throwing, which are not uniformly rubbish disposal. Person boxes lack duration; synthetic fallen/lying people lack verified public-sleep context. Independent human validation of historical assistant-reviewed unattended and sleeping labels remains missing. These sources are useful; their supported measurement differs from event validation.",
        warn=True,
    )

    tab(
        "Data / supervised partitions",
        "Training covered all five targets with uneven evidence",
        "First supervised study: 7,498 examples • 12,478 distinct view images",
        "data/clip_training_20261015/samples.jsonl + training findings",
        ["Objective", "Train (normal / +)", "Dev (normal / +)", "Test (normal / +)"],
        [
            ["Smoking, CCTV + train photos", "1,381 / 716", "722 / 413", "253 / 393"],
            ["Littering / dumping", "130 / 144", "29 / 32", "50 / 50"],
            ["Loitering", "282 / 41", "85 / 13", "64 / 11"],
            ["Public sleep", "49 / 25", "14 / 7", "12 / 6"],
            ["Unattended items", "416 / 74", "137 / 181", "720 / 242"],
        ],
        [285, 197, 197, 197],
        "A separate 806-photo smoking holdout is reported separately. Endpoints and adjacent frames remain correlated.",
        "Training totals 3,258, development 1,633, test 1,801, plus 806 Stanford40 test photographs (666 normal / 140 smoking). Classes are target-specific, with unknown labels for other objectives masked. CCTV camera c4 trains, c2 develops, c1 tests; original photos are included in training/development where partitions permit. Public sleep is provisional. Unattended endpoint counts represent only two held-out camera families, not 962 events. Later strategies expand train to 3,973 and dev to 1,812 with additional photos.",
        row_h=43,
        font=13,
    )

    tab(
        "Data / acquisition expansion",
        "More data was acquired and screened for compatibility",
        "Source diversity, duplicate control and label meaning matter more than download size",
        "Strategy acquisition receipts + pilot manifests",
        ["Source", "Acquired / retained", "Role and limit"],
        [
            [
                "Smoking photos + HMDB",
                "1,118 photos; 906 clips",
                "Photo learning; action-source transfer",
            ],
            [
                "Thermal loitering",
                "874,936 observations; 31 dates",
                "Provided tracks; one thermal camera",
            ],
            [
                "MEVA + public sleep",
                "13 videos; 28 new-author photos",
                "Action proxy; provisional photo labels",
            ],
            [
                "AVSS2007",
                "3 videos; 2 clean primary events",
                "One conservative source block",
            ],
            [
                "17 Oct author release",
                "3 new clips + 3 ABODA aliases",
                "30 highway-debris component frames",
            ],
        ],
        [210, 322, 344],
        "The latest release adds small-object stress cases, but adds no verified owner-associated unattended events.",
        "1,120 new smoking photos were downloaded and 1,118 retained; 715 train,179 dev,224 test. HMDB: 109 smoking and 797 confusable action clips before grouped splits. Thermal: 19,737 track IDs, 17,878 eligible complete trajectories, 19/6/6 chronological dates. MEVA: 15 action windows plus12 pre-event controls,10 synchronization groups across7 cameras, not persistent unattended truth. Public sleep:40 downloaded,29 usable,28 author-screened cases(17+/11normal;25 author groups). AVSS HARD is censored and excluded from the primary two-event result. New author clips show highway debris, not owner histories; three other entries explicitly reference ABODA and are not new sources.",
        row_h=43,
        font=13,
    )

    paired(
        "Methodology / generalization",
        "Keep related sources together and freeze decisions",
        "A source holdout is more informative than a random split of adjacent frames",
        "Input hashes; source guard; saved development selections",
        (
            "Before evaluation",
            "Preserve source / camera / author groups\nRespect official video/photo partitions\nRemove exact-pixel aliases\nPurge known shared-film part aliases\nFit routes and thresholds on dev only",
        ),
        (
            "Limits that remain",
            "Previously inspected tests are retrospective\nMovie / actor identities are incomplete\nOnly a few physical cameras / events\nProvisional interval / sleep labels\nRepeated seeds add no new domains",
        ),
        "No test-optimal threshold or lucky seed is substituted after seeing results.",
        "The actual fine-tuning primary is seed42;43/44 show optimization variability. The later video pilot retains all seeds and uses worst dev-source AP then meanAP; final selections precede scoring new AVSS. Its source guard removes10 known film-alias training clips and no test clips. Exact pixels are verified for25,452 pilot inputs; hashes and source-name grouping do not prove all near-duplicate/actor identities are known. On17Oct, AVSS and historical scores are already inspected, so all policy comparisons there are retrospective even though choices still use development scores.",
        warn=True,
    )

    paired(
        "Testing / compute and numerical controls",
        "Accuracy runs used available GPU compute",
        "No pacing, CPU emulation or new latency benchmarks in the later studies",
        "GPU state files + numerical / live-checkpoint audits",
        (
            "Quality execution",
            "RTX 5000 Ada GPU\nBatched CLIP extraction / head fitting\nDeterministic VLM generation\nFailures and abstentions stay counted\nPinned weights, inputs and raw scores",
        ),
        (
            "Numerical consistency",
            "Fine-tuning: BF16 autocast, F32 weights\nFinal CLIP scoring: F32, TF32 off\nBatch-shape sensitivity was audited\nNew detector: 830 F32 GPU views\nEvent delay measures task detection",
        ),
        "Event onset delay is a quality metric; it is separate from computation latency.",
        "Expanded image comparisons use shared letterbox preprocessing capped at768pixels, CLIP float32 and unquantized VLM automatic BF16,64-token deterministic generation, MiniCPM one-slice baseline. Those are explicit baseline settings, not a search of optimal native video methods. The fine-tuning audit found AMP inference sensitivity to batch shape; weights and routes were fixed, final evaluation rerunF32 and thresholds refitted usingF32dev only. Model bundles and feature/checkpoint replays record numerical tolerances. The isolated CUDA runtime was validated with actual computation and used for the new detector comparison.",
    )

    paired(
        "Evaluation / what each metric answers",
        "Ranking, frame decisions and alerts are different",
        "All three are needed to judge practical detection",
        "Binary metrics + strict causal event accounting",
        (
            "Image / clip discrimination",
            "AP and AUROC: score ranking\nRecall + FPR: a fixed operating point\nF1 + confusion counts: class balance\nVLM hard labels: no invented AP\nWeak candidates: prediction rate only",
        ),
        (
            "Continuous event detection",
            "Past-only 20-second histories\nActual alert emission times\nOne matched alarm per labelled event\nNormal false, repeated, unknown alerts\nOnset delay + reviewed normal hours",
        ),
        "An alert in an unreviewed transition is unknown; an early continuing alarm is not silently rematched to a later onset.",
        "Pilot global rule:2s persistence,1s release,10s cooldown, reset for score gaps>1.1s. Thresholds maximize dev event recall within the stated normal false-alert budget. Strict matching counts only alert starts within positive intervals and separately reports repeated event alarms. Historical onset windows are provisional, so some unknown-time alarms may be useful early detections; do not portray those as proven normal false alarms. Overlapping windows do not create independent event observations. Confidence from a few cameras is limited.",
    )

    raw = {
        r["model"]: r
        for r in quality["rows"]
        if r["cohort"] == "unattended_expanded_frame" and r["split"] == "test"
    }
    tab(
        "Baseline / four-model image comparison",
        "Expanded unattended images did not yield a clear winner",
        "962 test endpoints • 242 positive / 720 normal • two camera families",
        "Expanded accuracy quality_summary.json / raw frame decisions",
        ["Model", "Precision", "Recall", "F1", "FPR"],
        [
            [name]
            + [
                f"{raw[key][k]:.1%}" if k != "f1" else f"{raw[key][k]:.3f}"
                for k in ["precision", "recall", "f1", "fpr"]
            ]
            for key, name in zip(d.NAMES, d.SHORT)
        ],
        [292, 146, 146, 146, 146],
        "MiniCPM has the best raw frame F1 here, while missing 65% of positive endpoints. Labels remain provisional.",
        "These are 14 October results preceding CLIP adaptation. MiniCPM has 84 TP, 25 FP, 158 FN and 695 TN (frame F1 .479). With 20-second contact sheets recall falls to 10.74%, FPR .14% and F1 .193. SmolVLM has 1,457 abstentions over 13,789 evaluations; MiniCPM has 45; both have zero runtime errors. Hard categorical responses have no real continuous ranking scores. These models have not been rerun or fine-tuned in subsequent CLIP studies. A contact-sheet comparison is not native temporal-video inference.",
        row_h=52,
        font=14,
        warn=True,
    )

    paired(
        "Actual training / six visual fine-tunes",
        "Both CLIP visual backbones were updated",
        "Two encoders × seeds 42 / 43 / 44 • frozen early layers and text towers",
        "Fine-tuning changed-tensor audit + training protocol",
        (
            "What learned",
            "MobileCLIP: 18.1M visual parameters\nSigLIP: 14.2M visual parameters\nLast vision stage + five binary heads\nMasked target losses and balancing\nScene-preserving augmentation",
        ),
        (
            "How overfitting was constrained",
            "Pretrained feature / weight penalties\nFrozen early layers and batch norms\nProbe and pretrained-score controls\nDevelopment-selected routes / epochs\nThree-seed variation retained",
        ),
        "These six visual fine-tunes are distinct from later frozen-feature classifiers, GRUs and ensembles.",
        "The study has 7,498 target-specific examples and 12,478 distinct views. Temporal features use last/mean/max/delta pooling, preserving source groups. Probe C values 1/10/100 and fine-tuning settings are selected on dev. Seed 42 was primary before held-out results. Six checkpoints contain finite changed visual tensors, checked by delta norms and hashes. Verification replays 46,926 held-out and 29,394 dev predictions. Models classify scenes/clips and do not themselves output object boxes or ownership.",
    )

    tab(
        "Actual training / held-out ranking",
        "Fine-tuning helped some tasks and hurt others",
        "Same-study pretrained score → development-selected primary route (AP)",
        "CLIP fine-tuning findings / seed 42",
        ["Evidence", "Cases", "MobileCLIP AP", "SigLIP AP"],
        [
            ["Smoking CCTV camera c1", "646", ".642 → .564", ".669 → .650"],
            ["Waste dumping clips", "100", ".424 → .858", ".499 → .885"],
            ["Loitering windows", "75 (11 +)", ".320 → .093", ".097 → .097"],
            ["Public-sleep photos", "18 (6 +)", ".502 → .976", ".733 → .915"],
            ["Unattended endpoints", "962", ".209 → .652", ".309 → .896"],
        ],
        [304, 150, 211, 211],
        "Photo smoking improves strongly (.901 / .954 AP), but photo success did not transfer to CCTV.",
        "These are source holdouts that were historically inspected. Gains combine supervision and temporal representations; some routes retain a probe or pretrained score based on dev. On 806 photo tests SigLIP AP rises .499 to .954 and MobileCLIP .220 to .901. Loitering has 11 positives among 75 windows (prevalence .147), with both selected APs below prevalence. Sleep has six positive photos and provisional labels. Unattended labels are provisional and cover only two held-out camera families. All seeds and alternatives remain preserved.",
        row_h=43,
        font=14,
        warn=True,
    )

    paired(
        "Strategy exploration / 330 configurations",
        "The next screen combined data, crops and adaptation",
        "33 configurations per target / backbone • five objectives • two CLIP encoders",
        "strategy_study_20261016 / frozen development selections",
        (
            "Inputs and representations",
            "Train: 3,258 → 3,973 examples\nAutomatic contextual person/item crops\nFull / region visual representations\nTemporal pooling vs ordered views\nSparse detection geometry",
        ),
        (
            "Candidate routes",
            "Linear and nonlinear classifiers\nPretrained / adapted weight blends\nAlpha: 0, .25, .5 and 1\nDev-source AP route selection\nRaw, dev-F1 and FPR-budget cutoffs",
        ),
        "330 configurations include references and 320 classifier fits; they are not 330 separate backbone fine-tunes.",
        "Earlier adapted weights are reused; this is not a new visual fine-tune per configuration. Seed 42 is primary. Automatic crops and sparse geometry do not supply full MOT or ownership. Verification replays 117,612 predictions and checks 20,980 image identities and actual inference for all five tasks on both encoders.",
    )

    tab(
        "Strategy exploration / historical source diagnostics",
        "Historical AP rose unevenly after the strategy screen",
        "Previous selected bundle → new selected strategy • same previously inspected cases",
        "Strategy findings / primary seed 42",
        ["Objective", "MobileCLIP AP", "SigLIP AP", "Evidence limitation"],
        [
            ["Smoking CCTV", ".564 → .690", ".650 → .722", "One held-out camera"],
            [
                "Waste dumping",
                ".858 → .806",
                ".885 → .893",
                "Full-clip action classification",
            ],
            ["Loitering", ".093 → .489", ".097 → .550", "One Street Scene camera"],
            [
                "Public sleep",
                ".976 → .931",
                ".915 → .883",
                "Small provisional photo set",
            ],
            [
                "Unattended",
                ".652 → .980",
                ".896 → .925",
                "Provisional overlapping endpoints",
            ],
        ],
        [181, 201, 201, 293],
        "MobileCLIP dumping and both public-sleep photo routes regress. Higher AP still does not solve alert thresholds.",
        "These development-selected routes concern already inspected historical holdouts. Crops help smoking and geometry helps loitering, while dumping/public-sleep regressions remain. Dev-F1 FPR is still 75.9% / 84.6% on CCTV smoking and 99.7% / 100% on unattended items. SigLIP unattended zero-logit F1 falls from .749 to .131. Preserve threshold failures alongside AP gains.",
        row_h=43,
        font=13,
        warn=True,
    )

    tab(
        "Generalization / new-source transfer",
        "The broader transfer tests exposed remaining failures",
        "Previous bundle → selected strategy AP • before later HMDB adaptation",
        "Strategy transfer report + source-screening artifacts",
        ["Transfer evidence", "Cases", "MobileCLIP AP", "SigLIP AP"],
        [
            ["New smoking-photo holdout", "224", ".953 → .997", ".995 → 1.000"],
            ["HMDB movie/web action clips", "906", ".605 → .576", ".685 → .664"],
            ["New-author public-sleep photos", "28", ".925 → .952", ".943 → .933"],
            ["MEVA abandonment-action proxy", "27", ".569 → .545", ".637 → .621"],
        ],
        [344, 110, 211, 211],
        "Photo-domain gains coexist with movie/web and MEVA regressions. These sources are not interchangeable.",
        "The 224-photo holdout shares its collection with new training photos. HMDB was an unseen source here; both routes regress on 906 clips. Later HMDB adaptation uses 530 training and 205 test clips, a different protocol. Public sleep has 28 provisional new-author cases, 17 positive and 11 normal. MEVA has 15 action windows and 12 controls, not verified persistent unattended state.",
        row_h=50,
        font=14,
        warn=True,
    )

    paired(
        "All five objectives / temporal evidence",
        "Loitering and public sleep remain data-limited",
        "Their useful component results do not yet validate robust event detection",
        "Thermal trajectory benchmark + public-sleep source audit",
        (
            "Loitering component",
            "17,878 eligible supplied trajectories\nChronological 19 / 6 / 6-date split\nForest test AP .577 vs duration .283\nDevelopment AP .863\nOne thermal camera; no RGB tracking",
        ),
        (
            "Public-sleep component",
            "28 new-author photo cases\nMobileCLIP AP .952; SigLIP .933\nProvisional visible-sleep labels\nNo confirmed persistence / duration\nRest, lying and bedding need controls",
        ),
        "The next model must see the evidence the target requires: complete tracks or persistent state, with matched negatives.",
        "Thermal data supplies 874,936 observations and 19,737 track IDs; 17,878 eligible trajectories remain. A selected 400-tree depth-eight forest has test AP .577, F1 .664 and FPR .187 versus duration-only AP .283. Development AP .863 exposes date transfer loss. These are supplied tracks from one thermal camera. RGB Street Scene is separate and has only 11 positive test windows. Photos cannot establish sleep duration.",
        warn=True,
    )

    paired(
        "Video pilot / error audit",
        "Small people and missing items were concrete bottlenecks",
        "117 deliberately selected errors were provisionally reviewed across the five targets",
        "Pilot crop-coverage and CAVIAR component audits",
        (
            "Smoking spatial evidence",
            "Median smoker height: 11.2 px at 224\n71.8% of positive frames below 12 px\nOld crop misses half every target: 40.2%\nMulti-person / upper crops: 0.3%\nCoverage is separate from recognition",
        ),
        (
            "Item and ownership evidence",
            "CAVIAR item recall at IoU .5:\nFaster R-CNN tracks: 0 / 136\nGrounding DINO tracks: 66 / 136\nProximity can confuse passers-by\nJitter can reset stationarity",
        ),
        "Error-board findings are triage, not an independently reviewed estimate of failure causes.",
        "At 224 input, the median largest smoker is 11.2 pixels tall; all targets are below 12 pixels in 71.8% of positive CCTV frames. Automatic multiple-person/upper-body crops reduce losing half every target from 40.2% to 1/393. Detector resolution also changes; coverage is not recognition. CAVIAR partial item-reference recall rises from 0/136 to 66/136. Unmatched tracks are not proven false positives and ownership has no independent truth.",
    )

    s = new(
        "Video pilot / causal architecture",
        "Separate appearance from tracked event state",
        "Automatic proposals → frozen CLIP features → causal heads → explicit alerts",
        "video.py + temporal.py + training / inference scripts",
        "This pilot freezes pretrained encoders, separate from earlier visual fine-tuning. Dense 2 Hz video and 16 HMDB views provide scene, person, upper-body and item contexts. Eighteen geometry values come from causal Hungarian IoU/velocity tracking, not ByteTrack or ReID. GRUs use 128 projection and 64 hidden dimensions with past-only histories and gap timing. The cached-feature replay CLI is not a new-video streaming application.",
    )
    for x, title, body in [
        (
            42,
            "1. Detect",
            "Full scene + up to four persons\nWhole / upper-body crops\nTwo item / owner contexts",
        ),
        (
            340,
            "2. Represent",
            "Frozen CLIP visual features\nCausal geometry and timestamps\nLinear / GRU / instance heads",
        ),
        (
            638,
            "3. Decide",
            "Past-only scores / same-item state\nPersistence and alert emission\nEvent / normal / unknown counts",
        ),
    ]:
        d.card(s, x, 182, 280, 224, title, body, size=15)
    brief(
        s,
        "No annotation boxes choose inference crops. Missing owner history never becomes an invented ownership relation.",
    )

    paired(
        "Video pilot / controlled learning",
        "Fifty heads tested region, sequence and person evidence",
        "8,931 modeling examples after the known-film source guard",
        "Final joint_selection.json + saved training histories",
        (
            "Smoking: 14 heads / encoder",
            "Scene and region linear baselines\nRegion GRU, three seeds\nRegion + geometry GRU, three seeds\nPerson-supervised MIL, three seeds\nHard-negative variant, three seeds",
        ),
        (
            "Unattended: 11 heads / encoder",
            "Scene and region linear baselines\nRegion GRU, three seeds\nCOCO-geometry GRU, three seeds\nGrounding-geometry GRU, three seeds\nUnchanged visual CLIP features",
        ),
        "Route and epoch selection maximize worst development-source AP, then mean AP. All seeds and regressions remain visible.",
        "HMDB contributes 530 train, 161 dev and 205 test clips after ten known shared-film training aliases are purged. No test clip is removed. Person supervision uses 1,460 matched automatic training-camera regions with 277 positives. Training-only hard-negative mining starts at epoch four; checkpoints selected earlier do not demonstrate its benefit. Heads use float32 AdamW, learning rate .0005, at most 60 epochs and patience eight. Most of 936 prepared recordings are action clips, not cameras.",
    )

    tab(
        "Video pilot / smoking ranking",
        "CCTV gains did not repeat consistently across backbones",
        "205 HMDB evaluation clips now follow 530 grouped training clips; CCTV remains static",
        "Pilot report: same-case AP comparisons",
        ["Model / source", "Previous bundle", "Pilot scene", "Selected pilot"],
        [
            ["MobileCLIP / CCTV 646", ".690", ".626", ".725"],
            ["SigLIP / CCTV 646", ".722", ".717", ".688"],
            ["MobileCLIP / HMDB 205", ".697", ".842", ".908"],
            ["SigLIP / HMDB 205", ".762", ".888", ".778"],
        ],
        [354, 174, 174, 174],
        "SigLIP’s development-selected pilot route regresses on CCTV and trails its own scene baseline on HMDB.",
        "The previous bundle is recomputed on the exact same 205 HMDB cases, using eight views and no HMDB training versus 16 views and 530 training clips now. Gains combine data and methods. At the old pooled dev-budget point, CCTV recall/FPR are 16.5%/7.9% for MobileCLIP and 2.3%/1.6% for SigLIP. HMDB values are 86.7%/2.9% and 66.7%/1.7%. Movie/web recognition does not establish CCTV transfer.",
        row_h=50,
        font=14,
        warn=True,
    )

    rows = []
    for model in MODELS:
        for label, name in [
            ("previous", "previous_primary_original_policy"),
            ("17 Oct", selection[model]["selected"]),
        ]:
            v = smoke[model][name]["sources"]
            c, h = v["smoking_cctv"], v["hmdb51_action_transfer"]
            rows.append(
                [
                    SHORT[model] + " / " + label,
                    f"{c['ap']:.3f}",
                    f"{c['recall']:.1%} / {c['fpr']:.1%}",
                    f"{h['ap']:.3f}",
                    f"{h['recall']:.1%} / {h['fpr']:.1%}",
                ]
            )
    tab(
        "17 Oct improvement / smoking ensembles",
        "Ensembling improves HMDB; low-FPR CCTV remains hard",
        "38 dev-only policies • every dev source must meet the 5% empirical FPR constraint",
        "smoking_selection.json + smoking_results.json",
        ["Policy", "CCTV AP", "Recall / FPR", "HMDB AP", "Recall / FPR"],
        rows,
        [244, 116, 200, 116, 200],
        "SigLIP selects the all-head ensemble. MobileCLIP retains the same person head with a stricter cutoff.",
        "Each encoder compares 14 previous heads, four seed ensembles and one all-head ensemble: 19 policies. Equal sigmoid averages are not calibrated probabilities. Both selections freeze before held-out scores load. Every dev source must meet the FPR budget under one cutoff. SigLIP HMDB AP improves .7779 to .9645 with recall 66.7% to 76.7% at observed FPR 1.71%; CCTV recall remains 1.78%. MobileCLIP retains its head with a stricter cutoff. These are retrospective cohorts and reused development sources.",
        row_h=50,
        font=13,
        warn=True,
    )

    tab(
        "Video pilot / unattended events",
        "Explicit item state was more useful than a global score",
        "Two clean AVSS events • 2.76 normal minutes • first used 16 Oct, now retrospective",
        "Pilot events.csv / clean primary source subset",
        ["Route / policy", "Events", "Normal false / repeats", "Median onset delay"],
        [
            ["MobileCLIP / global score", "2 / 2", "1 / 3", "25.8 s"],
            ["SigLIP / global score", "2 / 2", "6 / 3", "12.1 s"],
            ["MobileCLIP + known state gate", "2 / 2", "0 / 0", "36.3 s"],
            ["SigLIP + known state gate", "2 / 2", "0 / 1", "34.8 s"],
            ["Known source state rule alone", "2 / 2", "0 / 0", "34.8 s"],
        ],
        [365, 95, 229, 187],
        "The rule requires the same stationary item for 30 s and its previously associated owner absent for ≥0.5 s.",
        "Global alerts use two-second persistence, one-second release and ten-second cooldown. The one-false-alert/hour development target has only 28 seconds of dev normal footage and cannot certify that rate. Observed global normal rates are 21.7/hour and 130.3/hour. The separate AVSS source rule requires the same item stationary for 30 seconds and its associated owner absent. State-only results are independent of CLIP. The three videos appear to share one platform view; HARD remains a censored diagnostic.",
        row_h=43,
        font=13,
        warn=True,
    )

    paired(
        "Thresholds / calibration and uncertainty",
        "A good AP or zero observed alarms is not a guarantee",
        "Insufficient normal footage is a measurable limitation",
        "Disjoint-recording calibration + event rate bounds",
        (
            "Historical calibration diagnostic",
            "Entire normal recordings held aside\nBoth arms exclude those recordings\nSigLIP: 0/4 → 1/4 events\nNormal false alerts: 4 → 1\nStill 26.3 false alerts / normal hour",
        ),
        (
            "Uncertainty in the AVSS result",
            "2 events and 165.74 normal seconds\n0 observed state-rule false alerts\n95% one-sided Poisson upper bound:\n65.1 false alerts / normal hour\nCamera transfer remains untested",
        ),
        "Even under a homogeneous Poisson assumption, about 3 zero-alert normal hours are needed for an upper bound below 1/hour.",
        "Entire ABODA classroom video7 and walkway video2 provide disjoint normal calibration and are excluded from both evaluation arms. Four events and 137 normal seconds remain. MobileCLIP detects 0/4 in both arms; SigLIP changes 0/4 to 1/4, with four normal false alerts reduced to one, still 26.3/hour. AVSS lacks disjoint normal calibration. Zero alarms over 165.74 normal seconds gives a one-sided 95% Poisson upper rate of 65.07/hour, assuming a homogeneous process. It is not a camera-transfer guarantee.",
        warn=True,
    )

    er = events["aggregate"]["avss_primary_retrospective"]
    rows = []
    for name, label in [
        ("original_tracks_global_debounce", "Original global state debounce"),
        ("original_tracks_item_memory", "Original tracks + item memory"),
        ("robust_tracks_global_debounce", "Median stationarity + global"),
        ("robust_tracks_item_memory", "Median stationarity + memory"),
    ]:
        v = er[name]
        rows.append(
            [
                label,
                f"{v['detected_events']}/{v['events']}",
                f"{v['normal_false_alerts']} / {v['duplicate_alerts']}",
                f"{v['median_delay_seconds']:.2f} s",
            ]
        )
    tab(
        "17 Oct improvement / persistent item state",
        "More state machinery did not improve the AVSS result",
        "Original detections and 30-second source rule are held fixed",
        "event_results.json + replayed event_traces.jsonl",
        ["Ablation", "Events", "Normal false / repeats", "Median onset delay"],
        rows,
        [355, 95, 239, 187],
        "Keep the simpler state baseline: smoothing adds 0.25 s; item memory increases repeats on the censored HARD diagnostic.",
        "The new rule uses three trailing centers and one second of sustained .015-frame displacement; pending motion blocks alerts. Item memory uses spatial one-to-one association and rearms on five seconds of owner return, three seconds of motion or 30 seconds absent. Synthetic tests cover simultaneous items and ID gaps without proving field benefits. AVSS stays 2/2 with zero false/repeated alerts; smoothing raises delay from 34.818 to 35.068 seconds. Memory raises HARD repeats from one to two. Historical intervals use different provisional definitions.",
        row_h=50,
        font=13,
        warn=True,
    )

    paired(
        "17 Oct data / compatibility review",
        "The new highway clips are component stress tests",
        "Pinned author release: three bundled clips, three explicit ABODA aliases",
        "Author acquisition receipt + component_protocol.json",
        (
            "Acquisition and preparation",
            "28.6 MB author archive\nThree bundled highway videos\n7.6–14.6 seconds each\nTen annotated positive frames / video\nActual FPS and timestamps checked",
        ),
        (
            "What cannot be inferred",
            "Debris boxes do not establish an owner\nNo 30-second stationary evidence\nNo complete reviewed normal intervals\nABODA links add no independent data\nNo new unattended-event accuracy",
        ),
        "The additional frames test detection of very small objects; they are not newly verified unattended events.",
        "Source: https://github.com/hangsuuuu/video-of-abandoned-object-detection at revision 966439fcac1103b392c49d4135f128860a39d1f8. The receipt records content hashes. Three bundled clips show highway debris; three entries explicitly alias ABODA. Actual FPS is 30, about 14.966 and 19.912, with 357, 219 and 151 decoded frames. Beyond-stream timestamps are exposed, not rescaled. Thirty annotated frames form a component stress test with no owner or complete event labels. Visual review is provisional assistant review.",
        warn=True,
    )

    rows = []
    for source, label in [
        ("caviar_training_component", "CAVIAR item source boxes"),
        ("author_highway_debris_component", "Highway-debris stress cases"),
    ]:
        for iou in [0.3, 0.5]:
            a = {
                r["mode"]: r
                for r in component["aggregate"]
                if r["source"] == source and r["iou"] == iou
            }
            rows.append(
                [
                    label,
                    str(iou),
                    f"{a['full']['matched']}/{a['full']['reference_boxes']} ({a['full']['recall']:.1%})",
                    f"{a['full_plus_tiles']['matched']}/{a['full_plus_tiles']['reference_boxes']} ({a['full_plus_tiles']['recall']:.1%})",
                ]
            )
    tab(
        "17 Oct improvement / small-item detector",
        "Fixed overlapping tiles test the small-object bottleneck",
        "166 source-annotated frames × five views = 830 new GPU image views",
        "component_results.json + raw / merged detector predictions",
        ["Reference source", "IoU", "Full image recall", "Full + four tiles recall"],
        rows,
        [340, 72, 232, 232],
        "Same weights, query and score thresholds. Partial reference boxes support recall, not object precision or event accuracy.",
        "Grounding DINO tiny uses the same pinned weights, generic query and thresholds, float32 GPU inference and TF32 off. Four fixed 60%-size tiles overlap by 20%; detections merge with full-image output using NMS .5. Source boxes never choose tiles. There are 136 CAVIAR training-source item references and 30 highway-debris references. Partial annotations support recall, not precision. Tiling does not overwrite pilot event tracks or heads.",
        row_h=50,
        font=13,
    )

    vm = vad["metrics"][vad["selected_branch"]]
    paired(
        "VadCLIP / separate supervised study",
        "VadCLIP remains a qualified smoking baseline",
        "Offline supervised adaptation • different source groups and supervision",
        "gpu_all_models_20261013/vadclip/adapter/eval.json",
        (
            "Measured test evidence",
            f"1,056 test frames\n560 positive / 496 normal\nAP {vm['aupr']:.3f}; AUROC {vm['auroc']:.3f}\nRecall {vm['recall']:.1%}; FPR {vm['false_positive_rate']:.1%}\nF1 {vm['f1']:.3f}, dev-frozen cutoff",
        ),
        (
            "Comparison boundaries",
            "9 training clips / 4 validation clips\nAdapter receives smoking supervision\nNot an unsupervised five-task model\nNo new VadCLIP fine-tune this round\nNot the same split as CLIP camera c1",
        ),
        "Avoid a single leaderboard that mixes this offline supervised protocol with image-model zero-shot or event-stream tests.",
        "Original VadCLIP dual_mean gives AUROC .5874, AP .63323, F1 .56475, recall .51786 and FPR .35685 with a development-frozen threshold. Training uses 1,910 frames in nine clips; validation uses 312 frames in four clips and a 12.642M-parameter adapter. This offline supervised smoking protocol has different source groups and supervision from image models. The constant visual branch remains a recorded failure. No new VadCLIP training was done this round.",
    )

    tab(
        "Current assessment / all five objectives",
        "No single model is ready to cover all five objectives",
        "Use task-specific candidate routes and measure the evidence they actually consume",
        "All completed reports; stage-specific qualifications apply",
        ["Objective", "Useful candidate evidence", "Main unresolved issue"],
        [
            [
                "Smoking",
                "Person crops; SigLIP ensemble on HMDB",
                "CCTV low-FPR recall; tiny smokers",
            ],
            [
                "Littering",
                "SigLIP IWDD action AP .893",
                "Pedestrian litter vs placing; localization",
            ],
            [
                "Loitering",
                "Geometry / thermal supplied tracks",
                "Automatic RGB tracks; unseen cameras",
            ],
            [
                "Public sleep",
                "Photo probes / apparent posture",
                "Verified persistent video; rest controls",
            ],
            [
                "Unattended",
                "Grounded items + explicit state rule",
                "Owner truth; detection gaps; normal hours",
            ],
        ],
        [168, 343, 365],
        "These are research candidates with different protocols; the AP values do not form a comparable deployment leaderboard.",
        "SigLIP ensemble HMDB AP .9645 concerns 205 retrospective adapted movie/web clips; CCTV recall remains 1.8% at 1.6% FPR. MobileCLIP CCTV AP .725 does not ensure transferable thresholds. Dumping AP .893 uses 100 full-clip IWDD tests, not online localization. Street Scene and thermal tracks use different protocols/sensors. Public-sleep photos remain provisional; unattended evidence has two primary events and a few normal minutes. Each candidate needs target-camera event validation.",
        row_h=43,
        font=13,
        warn=True,
    )

    paired(
        "What worked / what did not",
        "Retain demonstrated gains and preserve regressions",
        "Model selection must reward usable decisions as well as ranking",
        "All saved source-level metrics and candidate histories",
        (
            "Useful progress",
            "Real dumping-action supervision\nSource-grouped data / duplicate checks\nPerson-region coverage for smoking\nExplicit item persistence / owner state\nEnsembles on adapted action clips",
        ),
        (
            "Not established or regressed",
            "Photo gains → surveillance transfer\nAP gains → stable alert thresholds\nExtra GRU / mining → universal gains\nMore memory → fewer real alarms\nMore adjacent frames → more events",
        ),
        "The latest event-policy ablations remain research alternatives; the simpler baseline is retained.",
        "Preserved regressions include CCTV fine-tuning, MobileCLIP dumping, sleep photos, HMDB/MEVA transfer, SigLIP pilot CCTV, conservative-cutoff recall loss and HARD memory repeats. Hard-negative runs can select before mining begins. Repeated development searches can overfit sources even when no test selection occurs. More adjacent frames, seeds and hyperparameter trials do not add independent cameras/events. Tiling is a component finding, not an established event-pipeline improvement.",
        warn=True,
    )

    tab(
        "Next steps / generalization gates",
        "Prioritize independently labelled target-camera evidence",
        "Freeze the next source holdout before looking at its predictions",
        "Existing definitions; current component and event bottlenecks",
        ["Priority", "Concrete work", "Acceptance evidence"],
        [
            [
                "1. Data / labels",
                "Review ownership, event boundaries, controls",
                "Independent cameras / events, not frame count",
            ],
            [
                "2. Components",
                "Validate crops, item recall and associations",
                "No oracle crops; identity / coverage audits",
            ],
            [
                "3. Training",
                "Train on confirmed source-balanced events",
                "Fixed dev operating point; preserve controls",
            ],
            [
                "4. Final evaluation",
                "New untouched cameras + long normal streams",
                "Event recall, delay, duplicates and FAR bounds",
            ],
        ],
        [145, 366, 365],
        "Preserve an abstain / unknown state when duration, ownership or visibility is insufficient.",
        "Acquire compatible evidence with matched controls: smoking versus drinking/eating/coughing; disposal versus benign placing/carrying; complete dwell tracks; persistent apparent sleep versus rest/lying; owner-associated items with retrieval and occlusion controls. Independent target-camera review and event diversity are still needed. Freeze a fresh camera holdout before scores are inspected. Normal-hour bounds depend on assumptions and do not certify source transfer. Missing duration, ownership or visibility should permit unknown decisions.",
        row_h=50,
        font=13,
    )

    paired(
        "Reproducibility / artifacts and checks",
        "Every stage remains inspectable and replayable",
        "Editable PowerPoint and vector PDF share one checked layout",
        "Progress source manifest + verification / validation JSON",
        (
            "Preserved research evidence",
            f"Fine-tuning: {fine['predictions_recomputed']:,} held-out scores\nStrategy screen: {strategy['predictions_recomputed']:,} scores\nVideo pilot: {pilot['predictions_replayed']:,} scores\nLatest: {latest['smoking_predictions_replayed']:,} smoking scores\nRaw detector views + causal alert traces",
        ),
        (
            "Validation and deliverables",
            f"{validation['unit_tests']} unit tests passed\nRuff lint / format + whitespace checks\n{latest['development_policies_checked']} policy selections checked\n{latest['component_metric_cells_recomputed']} detector-metric cells replayed\nPDF / PPTX + notes + source hashes",
        ),
        "Historical artifacts remain unchanged. New results, qualifications and code are linked from the latest report.",
        "Latest report: results/event_improvements_20261017/report.md; overview: reports/event_improvements_findings_20261017.md. Earlier expanded, fine-tuning, strategy and video-pilot reports remain preserved. New scripts replay ensemble selection, causal memory and detector merges/matches; source/code/weight hashes are retained. The slide builder reads saved results and runs no inference. Editable PPTX and vector PDF share one checked layout, with notes and source hashes. Earlier decks remain historical.",
    )

    assert len(d.SLIDES) == 32, len(d.SLIDES)
    assert (
        component_protocol["reference_boxes"]["author_highway_debris_component"] == 30
    )


def main():
    build_content()
    d.check_geometry(d.SLIDES)
    for stem, slides in [
        (STEM, d.SLIDES),
        ("summary_slide_progress_20261017", [deepcopy(d.SLIDES[0])]),
    ]:
        path = OUT / (stem + ".pptx")
        d.render_pptx(slides, path)
        ppt = Presentation(path)
        assert len(ppt.slides) == len(slides)
        ppt.core_properties.title = "Five-anomaly evaluation and training progress"
        ppt.core_properties.subject = (
            "Methodology, data, findings and improvements through 17 October 2026"
        )
        ppt.save(path)
        d.render_pdf(slides, OUT / (stem + ".pdf"))
    notes = [
        "# Evaluation and training progress — speaker notes",
        "",
        "Finalized 17 October 2026 (Asia/Jakarta). Historical protocols stay separate.",
        "",
    ]
    for i, s in enumerate(d.SLIDES, 1):
        notes += [f"## {i:02d}. {s['title']}", "", s["notes"], ""]
    (OUT / "progress_speaker_notes_20261017.md").write_text("\n".join(notes))
    (OUT / "progress_source_manifest_20261017.json").write_text(
        json.dumps(
            {
                "date": "2026-10-17",
                "timezone": "Asia/Jakarta",
                "slide_count": len(d.SLIDES),
                "sources_sha256": d.SOURCES,
                "builders_sha256": {
                    p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                    for p in [Path(__file__), OUT / "build_presentation.py"]
                },
                "layout": d.SLIDES,
                "pdf_rendering": "Independent vector PDF from same measured layout; embedded fonts.",
            },
            indent=2,
        )
        + "\n"
    )
    print(
        "CREATED", len(d.SLIDES), "slides, PPTX/PDF, summary, notes and source manifest"
    )


if __name__ == "__main__":
    main()
